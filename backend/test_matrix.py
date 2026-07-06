"""GLOBAL TEST MATRIX — the cross-feature "did we break anything" checklist.

One file that exercises the WHOLE system (reader → Brain pipeline → executor → serializer
→ API) across every row of the matrix, deterministically (LLM mocked where the Brain's
plan is needed; real multilingual *comprehension* is covered by the live suites).

  1  File types          .xlsx/.xlsm/.csv work; others rejected kindly
  2  Empty/huge/corrupt   handled without crashing
  3  Languages            operations work from Hindi/English/Urdu/mixed instructions
  4  Special characters   non-English text, emojis, symbols preserved (data + headers)
  5  Ambiguity            vague → a question, never a guess
  6  Unsupported          friendly "not yet", never a fabricated result
  7  Multi-step           sequences run in order; partial failures reported clearly
  8  Output integrity     valid .xlsx (opens in Excel/Sheets); live formulas; original untouched
  9  Errors               no stack trace / raw error ever shown
  10 Numbers vs text      numeric ops never silently treat numbers as text or vice-versa
  11 Empty results        "0 matched / nothing to do" stated clearly, not as an error

Run from backend:  .venv\\Scripts\\python.exe test_matrix.py
"""
from __future__ import annotations

import io
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import openpyxl
import pandas as pd
from fastapi.testclient import TestClient

from app import main
from app.executor import OperationError, execute_plan
from app.main import _serialize

passed = failed = 0


def check(name, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


_TECH = ("traceback", "exception", "keyerror", "valueerror", "0x", "line ", "/app/", "\\app\\",
         "pandas", "numpy", "nonetype")


def technical(msg):
    m = (msg or "").lower()
    return any(t in m for t in _TECH)


def run(ops, df, sheets=None):
    return execute_plan(df, ops, sheets)


client = TestClient(main.app, raise_server_exceptions=False)
_orig = main.llm.parse_instruction


def xlsx(rows, ext="xlsx"):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


print("GLOBAL TEST MATRIX\n")

# =========================================================================
# 1  File types
# =========================================================================
print("1  File types")
xb = xlsx([["A", "B"], [1, "x"], [2, "y"]])
r = client.post("/inspect", data={"session_id": "m_xlsx"}, files=[("files", ("d.xlsx", xb, "application/octet-stream"))]).json()
check("1 .xlsx loads", r["status"] == "ok" and r["tables"][0]["columns"][0]["name"] == "A", str(r)[:120])
r = client.post("/inspect", data={"session_id": "m_xlsm"}, files=[("files", ("d.xlsm", xb, "application/octet-stream"))]).json()
check("1 .xlsm loads", r["status"] == "ok", str(r)[:120])
r = client.post("/inspect", data={"session_id": "m_csv"}, files=[("files", ("d.csv", b"A,B\n1,x\n2,y\n", "text/csv"))]).json()
check("1 .csv loads", r["status"] == "ok", str(r)[:120])
r = client.post("/inspect", data={"session_id": "m_txt"}, files=[("files", ("notes.txt", b"hello", "text/plain"))]).json()
check("1 unsupported type rejected kindly", r["status"] == "error" and not technical(r["error"]) and "Excel" in r["error"], str(r)[:140])

# =========================================================================
# 2  Empty / huge / corrupt
# =========================================================================
print("\n2  Empty / huge / corrupt")
r = client.post("/inspect", data={"session_id": "m_empty"}, files=[("files", ("e.csv", b"A,B\n", "text/csv"))]).json()
check("2 empty (header-only) file → 0 rows, no crash", r["status"] == "ok" and r["tables"][0]["row_count"] == 0, str(r)[:120])
r = client.post("/inspect", data={"session_id": "m_corrupt"}, files=[("files", ("broken.xlsx", b"not a real xlsx at all", "application/octet-stream"))]).json()
check("2 corrupt file → friendly error, no crash", r["status"] == "error" and not technical(r["error"]), str(r)[:140])
huge = pd.DataFrame({"A": range(60000), "B": range(60000)})
main._remember_session("m_huge", {"tables": {"H": huge}, "primary": "H", "exts": {"H": "csv"}, "notes": {}})
r = client.post("/execute", data={"session_id": "m_huge", "plan": '{"operations":[{"action":"remove_duplicates"}]}'}).json()
check("2 huge file processes + shows a notice", r["status"] == "ok" and "large file" in " ".join(r["notes"]).lower(), str(r)[:120])

# =========================================================================
# 3  Languages — instructions in Hindi / English / Urdu / mixed flow through intact
# =========================================================================
print("\n3  Languages")
client.post("/inspect", data={"session_id": "m_lang"}, files=[("files", ("s.csv", b"Region,Revenue\nN,100\nS,200\n", "text/csv"))])
seen = {}
try:
    def spy(instruction, structure, history=""):
        seen["instr"] = instruction
        return {"operations": [{"action": "sort", "columns": ["Revenue"], "orders": ["desc"]}],
                "title": "Sort", "translation": "Sort by Revenue", "confidence": 90}
    main.llm.parse_instruction = spy
    for label, instr in [("Hindi", "रेवेन्यू के हिसाब से सॉर्ट करो"),
                          ("Urdu", "ریونیو کے حساب سے ترتیب دیں"),
                          ("Hinglish", "revenue ke hisaab se sort kar do")]:
        rr = client.post("/process", data={"instruction": instr, "session_id": "m_lang", "rewind": "-1", "history": ""}).json()
        check(f"3 {label} instruction processes end-to-end", rr.get("status") == "ok", str(rr)[:100])
        check(f"3 {label} instruction reached the Brain intact (UTF-8)", seen.get("instr") == instr, repr(seen.get("instr")))
finally:
    main.llm.parse_instruction = _orig

# =========================================================================
# 4  Special characters — emojis / symbols / non-English in data + headers
# =========================================================================
print("\n4  Special characters")
csv = "नाम 🙂,Price ₹,Tag©\nप्रिया,₹1,200,A—B\n😀Bob,€999,x✓\nप्रिया,₹1,200,A—B\n".encode("utf-8")
# (commas inside values would break naive CSV; use a clean one instead)
csv = "Name,Emoji,Symbol\nAnjali,😀,₹100\nBob,🎉,€50\nAnjali,😀,₹100\n".encode("utf-8")
client.post("/inspect", data={"session_id": "m_special"}, files=[("files", ("u.csv", csv, "text/csv"))])
r = client.post("/execute", data={"session_id": "m_special", "plan": '{"operations":[{"action":"remove_duplicates"}]}'}).json()
sample = r["preview"][0]["sample_rows"]
emojis = [row.get("Emoji") for row in sample]
check("4 emojis preserved in data", "😀" in emojis and "🎉" in emojis, str(emojis))
check("4 currency symbols preserved", any("₹" in str(row.get("Symbol")) or "€" in str(row.get("Symbol")) for row in sample), str(sample)[:120])
check("4 dedupe still works with emoji rows (3→2)", r["row_count"] == 2, str(r["row_count"]))
# header with an emoji round-trips
csv2 = "Region 🌍,Value\nN,1\nS,2\n".encode("utf-8")
r2 = client.post("/inspect", data={"session_id": "m_special2"}, files=[("files", ("h.csv", csv2, "text/csv"))]).json()
check("4 emoji/symbol header preserved", r2["tables"][0]["columns"][0]["name"] == "Region 🌍", str(r2["tables"][0]["columns"][0]))

# =========================================================================
# 5  Ambiguity → a question, never a guess
# =========================================================================
print("\n5  Ambiguity")
client.post("/inspect", data={"session_id": "m_amb"}, files=[("files", ("s.csv", b"Region,Revenue\nN,100\n", "text/csv"))])
try:
    main.llm.parse_instruction = lambda i, s, h: {"operations": [], "clarification": "Which column should I sort by — Region or Revenue?"}
    r = client.post("/parse", data={"instruction": "sort it", "session_id": "m_amb"}).json()
    check("5 vague instruction → clarify (no ops)", r["status"] == "clarify" and r["clarification"].endswith("?"), str(r)[:140])
finally:
    main.llm.parse_instruction = _orig

# =========================================================================
# 6  Unsupported → friendly "not yet", never fabricated
# =========================================================================
print("\n6  Unsupported")
try:
    main.llm.parse_instruction = lambda i, s, h: {"operations": [], "reply": "I can't do that yet — but I can sort, filter, dedupe, aggregate, chart, and more."}
    r = client.post("/parse", data={"instruction": "email this to my boss", "session_id": "m_amb"}).json()
    check("6 unsupported → friendly message, no ops", r["status"] == "message" and "can" in r["message"].lower(), str(r)[:140])
finally:
    main.llm.parse_instruction = _orig

# =========================================================================
# 7  Multi-step — order + clear partial failure
# =========================================================================
print("\n7  Multi-step")
client.post("/inspect", data={"session_id": "m_ms"}, files=[("files", ("s.csv", b"R,P\nN,3\nS,1\nN,2\n", "text/csv"))])
ok = client.post("/execute", data={"session_id": "m_ms",
    "plan": '{"operations":[{"action":"filter","conditions":[{"column":"R","operator":"equals","value":"N"}]},{"action":"sort","columns":["P"],"orders":["desc"]}]}'}).json()
check("7 steps run in order (filter then sort)", ok["status"] == "ok" and [row["P"] for row in ok["preview"][0]["sample_rows"]] == [3, 2], str(ok)[:140])
part = client.post("/execute", data={"session_id": "m_ms",
    "plan": '{"operations":[{"action":"sort","columns":["P"],"orders":["desc"]},{"action":"drop_columns","columns":["ghost"]}]}', "rewind": "0"}).json()
check("7 partial failure reported clearly", part["status"] == "ok" and part["partial"] is True and part["failed_step"] == 2 and "Step 2" in part["warning"], str(part)[:160])

# =========================================================================
# 8  Output integrity — valid xlsx, live formulas, original untouched
# =========================================================================
print("\n8  Output integrity")
fdf, _, render = run([{"action": "add_formula_column", "name": "Total", "formula": "{q} * {p}"}],
                     pd.DataFrame({"q": [2, 3], "p": [10, 10]}))
xbytes = _serialize(fdf, "out", "xlsx", render)[0]
wb = openpyxl.load_workbook(io.BytesIO(xbytes))  # opens cleanly == valid OOXML
cell = wb.active.cell(row=2, column=3).value
check("8 result is a valid .xlsx (opens cleanly)", wb.active.max_row == 3, f"rows={wb.active.max_row}")
check("8 formulas are LIVE (=A2*B2), not just values", isinstance(cell, str) and cell.startswith("=") and "A2" in cell and "B2" in cell, repr(cell))
# original untouched after an op via the API
client.post("/inspect", data={"session_id": "m_out"}, files=[("files", ("s.csv", b"R,P\nN,3\nS,1\nN,2\n", "text/csv"))])
before = main._SESSIONS["m_out"]["states"][0]["tables"]["s"].copy(deep=True)
client.post("/execute", data={"session_id": "m_out", "plan": '{"operations":[{"action":"remove_duplicates"}]}'})
after0 = main._SESSIONS["m_out"]["states"][0]["tables"]["s"]
check("8 original (state[0]) untouched after an operation", after0.equals(before), "original mutated!")

# =========================================================================
# 9  Errors — no stack trace / raw error ever shown
# =========================================================================
print("\n9  Errors")
_st = main.summarize_tables
try:
    main.summarize_tables = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("internal 0xBADF00D at line 7"))
    client.post("/inspect", data={"session_id": "m_err"}, files=[("files", ("s.csv", b"A\n1\n", "text/csv"))])
    r = client.post("/parse", data={"instruction": "sort by A", "session_id": "m_err"})
    b = r.json()
    check("9 unexpected crash → friendly error, no traceback", r.status_code == 500 and not technical(b.get("error", "")), str(b)[:140])
    check("9 internal details never leak", "0xBADF00D" not in str(b) and "RuntimeError" not in str(b), str(b)[:140])
finally:
    main.summarize_tables = _st

# =========================================================================
# 10  Numbers vs text — never silently confused
# =========================================================================
print("\n10  Numbers vs text")
# numbers stored AS TEXT sort numerically (1,2,10 — not 1,10,2)
out, _, _ = run([{"action": "sort", "columns": ["n"]}], pd.DataFrame({"n": ["10", "2", "1"]}))
check("10 numbers-as-text sort numerically", list(out["n"]) == ["1", "2", "10"], str(list(out["n"])))
# numeric filter on numbers-as-text compares numerically (>5 keeps 10, not '10'<'5')
out, _, _ = run([{"action": "filter", "conditions": [{"column": "n", "operator": "greater_than", "value": "5"}]}],
                pd.DataFrame({"n": ["10", "2", "1"]}))
check("10 numeric filter on text-numbers is numeric", list(out["n"]) == ["10"], str(list(out["n"])))
# summing a genuinely TEXT column is refused (never silently 0)
try:
    run([{"action": "aggregate", "agg_func": "sum", "agg_column": "Note"}], pd.DataFrame({"Note": ["a", "b"]}))
    check("10 sum of text column refused (not silent 0)", False, "no error")
except OperationError as e:
    check("10 sum of text column refused (not silent 0)", not technical(str(e)), str(e))
# a numeric formula on a text column is caught (self-correction explains)
try:
    run([{"action": "add_formula_column", "name": "X", "formula": "{Name} * 2"}], pd.DataFrame({"Name": ["a", "b"]}))
    check("10 numeric formula on text refused", False, "no error")
except OperationError as e:
    check("10 numeric formula on text refused", "numbers" in str(e).lower(), str(e))

# =========================================================================
# 11  Empty results — stated clearly, not an error
# =========================================================================
print("\n11  Empty results")
client.post("/inspect", data={"session_id": "m_zero"}, files=[("files", ("s.csv", b"Region,Revenue\nN,100\nS,200\n", "text/csv"))])
r = client.post("/execute", data={"session_id": "m_zero",
    "plan": '{"operations":[{"action":"filter","conditions":[{"column":"Region","operator":"equals","value":"Mars"}]}]}'}).json()
check("11 filter matching nothing → status ok (not an error)", r["status"] == "ok", str(r)[:120])
check("11 zero matches stated clearly with a count", r["row_count"] == 0 and ("0" in " ".join(r["notes"]) or "no" in " ".join(r["notes"]).lower()), str(r["notes"]))
# find/replace with nothing to change says so plainly
out, notes, _ = run([{"action": "find_replace", "column": "c", "find": "zzz", "replace": "q"}], pd.DataFrame({"c": ["a", "b"]}))
check("11 find/replace nothing → 'No cells' note, not an error", "no cells" in notes[0].lower(), notes[0])

main._SESSIONS.clear()
print(f"\n{passed} passed, {failed} failed.")
raise SystemExit(1 if failed else 0)
