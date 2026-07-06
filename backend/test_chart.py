"""Tests for the chart operation (Phase 2.2 — automatic chart generation).

Covers: correct type honored/defaulted, chart reflects current data, unsupported
chart explained, empty/one-row handled, and the chart is really embedded in the .xlsx.
Standalone script — run:  python test_chart.py
"""
import io
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook

from app.executor import execute_multi, OperationError
from app.operations.chart import chart
import app.main as m

passed = failed = 0


def check(name, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


print("Running chart checks...\n")

df = pd.DataFrame({"Month": ["Jan", "Feb", "Mar"], "Revenue": [100, 200, 150]})

# --- happy path ---
note, d = chart(df, {"action": "chart", "chart_type": "bar",
                     "x_column": "Month", "y_columns": ["Revenue"],
                     "chart_title": "Revenue by Month"})
check("bar chart directive built",
      d["type"] == "chart" and d["chart_type"] == "bar"
      and d["x_column"] == "Month" and d["y_columns"] == ["Revenue"], str(d))
check("chart reflects row count", d["rows"] == 3, str(d))
check("chart note is plain language", "bar chart" in note.lower(), note)

# correct type chosen: defaults to bar, and honors an explicit type
_, d_def = chart(df, {"action": "chart", "x_column": "Month", "y_columns": ["Revenue"]})
check("defaults to bar when type omitted", d_def["chart_type"] == "bar", str(d_def))
_, d_line = chart(df, {"action": "chart", "chart_type": "line",
                       "x_column": "Month", "y_columns": ["Revenue"]})
check("line chart type honored", d_line["chart_type"] == "line", str(d_line))

# pie shows a single series
df2 = pd.DataFrame({"Cat": ["a", "b"], "A": [1, 2], "B": [3, 4]})
_, d_pie = chart(df2, {"action": "chart", "chart_type": "pie",
                       "x_column": "Cat", "y_columns": ["A", "B"]})
check("pie uses a single series", d_pie["y_columns"] == ["A"], str(d_pie))

# --- failures (explained, not silent) ---
try:
    chart(df, {"action": "chart", "chart_type": "radar",
               "x_column": "Month", "y_columns": ["Revenue"]})
    check("unsupported chart type explained", False, "no error")
except OperationError as e:
    check("unsupported chart type explained", "radar" in str(e), str(e))

try:
    chart(df, {"action": "chart", "chart_type": "bar",
               "x_column": "Nope", "y_columns": ["Revenue"]})
    check("missing x column errors", False, "no error")
except OperationError:
    check("missing x column errors", True)

try:
    chart(df, {"action": "chart", "chart_type": "bar",
               "x_column": "Month", "y_columns": ["Month"]})
    check("non-numeric value column errors", False, "no error")
except OperationError:
    check("non-numeric value column errors", True)

# --- edges: empty / one row ---
try:
    chart(df.iloc[0:0], {"action": "chart", "chart_type": "bar",
                         "x_column": "Month", "y_columns": ["Revenue"]})
    check("empty data errors", False, "no error")
except OperationError:
    check("empty data errors", True)

_, d_one = chart(df.iloc[:1], {"action": "chart", "chart_type": "bar",
                               "x_column": "Month", "y_columns": ["Revenue"]})
check("one-row chart builds", d_one["rows"] == 1, str(d_one))

# --- integration: flows through the executor, data unchanged ---
res, name, notes, render = execute_multi({"t": df}, "t",
    [{"action": "chart", "chart_type": "bar", "x_column": "Month", "y_columns": ["Revenue"]}])
check("chart op flows into render directives",
      any(r.get("type") == "chart" for r in render), str(render))
check("chart op leaves the data unchanged",
      len(res) == 3 and list(res.columns) == ["Month", "Revenue"], str(list(res.columns)))

# chart reflects CURRENT data: aggregate (sum by Month) THEN chart -> 2 categories.
# aggregate renames the value column to sum_of_Revenue, so the chart targets that.
many = pd.DataFrame({"Month": ["Jan", "Jan", "Feb"], "Revenue": [100, 50, 200]})
res2, _, _, render2 = execute_multi({"t": many}, "t", [
    {"action": "aggregate", "agg_func": "sum", "agg_column": "Revenue", "group_by": ["Month"]},
    {"action": "chart", "chart_type": "bar", "x_column": "Month", "y_columns": ["sum_of_Revenue"]},
])
chart_dir = next(r for r in render2 if r.get("type") == "chart")
check("chart reflects aggregated data (2 months)", chart_dir["rows"] == 2 and len(res2) == 2,
      f"rows={chart_dir['rows']}, df={len(res2)}")

# --- the chart is really embedded in the worksheet / .xlsx ---
wb = Workbook()
ws = wb.active
ws.append(list(df.columns))
for row in df.itertuples(index=False):
    ws.append(list(row))
m._apply_chart(ws, df, {"type": "chart", "chart_type": "bar",
                        "x_column": "Month", "y_columns": ["Revenue"], "title": "Rev", "rows": 3})
check("chart added to the worksheet", len(ws._charts) == 1, str(len(ws._charts)))

# end-to-end serialize: CSV input upgrades to xlsx, opens cleanly, contains a chart part
out, fname, media = m._serialize(
    df, "sales.csv", "xlsx",
    [{"type": "chart", "chart_type": "bar", "x_column": "Month", "y_columns": ["Revenue"],
      "title": "Rev", "rows": 3}],
)
check("serialized file is an .xlsx", fname.endswith(".xlsx"), fname)
check("serialized xlsx opens cleanly", load_workbook(io.BytesIO(out)).active.max_row == 4)
names = zipfile.ZipFile(io.BytesIO(out)).namelist()
check("xlsx file actually contains a chart part",
      any("chart" in n for n in names), str(names))

print(f"\n{passed} passed, {failed} failed.")
raise SystemExit(1 if failed else 0)
