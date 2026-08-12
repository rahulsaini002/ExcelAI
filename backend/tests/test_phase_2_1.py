"""ENGINE PHASE 2.1 — pivot summaries (NO AI).

1-D and 2-D grids, totals computed from the RAW data (an averages Total is the true
overall average, not an average of averages), % of total (grand/row/column), date
bucketing on messy dates (real cells + three string formats + garbage), blank-group
honesty, the live =GROUPBY()/=PIVOTBY() mode with its Microsoft-365 warning, the
plan validator, and the failure paths.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_2_1.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from datetime import date
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p21.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app, _missing_columns  # noqa: E402
from app.operations.base import OperationError  # noqa: E402
from app.operations.pivot_summary import pivot_summary  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"
WB_PATH = TESTS / "standard_test_workbook.xlsx"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


# ---- deterministic messy fixture: real+string+garbage dates, text numbers, blanks ----
def make_fixture() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Date", "Region", "Product", "Price"])
    ws.append([date(2026, 1, 5), "North", "A", 100])
    ws.append(["2026-01-20", "North", "B", 50])
    ws.append(["Feb 10, 2026", "South", "A", "200"])   # text number
    ws.append(["15/03/2026", "South", "B", None])      # blank price
    ws.append([None, "South", "A", 150])               # blank date
    ws.append(["notadate", "North", "B", "abc"])       # garbage date + non-numeric
    ws.append(["2026-04-02", "South", "A", 40])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


FIX = make_fixture()
# ground truth: North=100+50=150, South=200+150+40=390, grand=540, 5 numeric cells


def run_ops(ops: list[dict], data: bytes = FIX, fname: str = "pivot.xlsx"):
    sid = f"p21-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", (fname, data, OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": ops})})
    j = r.json()
    book = grid = None
    if j.get("status") == "ok" and j.get("download_id"):
        raw = client.get(f"/download/{j['download_id']}").content
        book = openpyxl.load_workbook(io.BytesIO(raw))
        grid = pd.read_excel(io.BytesIO(raw))
    return j, book, grid


def row(grid: pd.DataFrame, col: str, label: str) -> pd.Series | None:
    hit = grid[grid[col].astype(str) == label]
    return hit.iloc[0] if len(hit) else None


def text(j: dict) -> str:
    return " ".join([j.get("explanation") or "", " ".join(j.get("notes") or [])])


print("ENGINE PHASE 2.1 — pivot summaries (no AI)\n")

# ---- (a) 1-D summary + totals ---------------------------------------------------------
j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                       "value_column": "Price"}])
ok = j.get("status") == "ok" and grid is not None
check("1-D: status ok + grid downloads", ok, str(j)[:200])
if ok:
    check("1-D: per-group sums right (North 150, South 390)",
          row(grid, "Region", "North")["sum_of_Price"] == 150
          and row(grid, "Region", "South")["sum_of_Price"] == 390, grid.to_string())
    check("1-D: Total row = raw grand total 540",
          row(grid, "Region", "Total") is not None
          and row(grid, "Region", "Total")["sum_of_Price"] == 540, grid.to_string())
    check("1-D: blanks/non-numeric honestly noted",
          "1 blank cell" in text(j) and "1 non-numeric cell" in text(j), text(j)[:300])

# ---- (b) 2-D summary + totals row AND column ------------------------------------------
j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                       "pivot_column": "Product", "value_column": "Price"}])
ok = j.get("status") == "ok" and grid is not None
check("2-D: status ok", ok, str(j)[:200])
if ok:
    n = row(grid, "Region", "North")
    t = row(grid, "Region", "Total")
    check("2-D: cross cells right (North: A 100, B 50)",
          n is not None and n["A"] == 100 and n["B"] == 50, grid.to_string())
    check("2-D: totals column and totals row (North 150 · A 490 · grand 540)",
          n["Total"] == 150 and t is not None and t["A"] == 490 and t["Total"] == 540,
          grid.to_string())

# ---- (c) totals are computed from RAW data (average ≠ average of averages) ------------
j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                       "value_column": "Price", "agg_func": "average"}])
if grid is not None:
    t = row(grid, "Region", "Total")
    check("average Total = true overall average 108 (not mean-of-means 102.5)",
          t is not None and abs(t["average_of_Price"] - 108) < 1e-9, grid.to_string())
    check("average note explains the Total line honestly",
          "not an average of averages" in text(j), text(j)[:300])
else:
    check("average pivot ran", False, str(j)[:200])

# ---- (d) % of total: grand and row ----------------------------------------------------
j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                       "pivot_column": "Product", "value_column": "Price",
                       "percent_of": "grand"}])
if grid is not None:
    n, t = row(grid, "Region", "North"), row(grid, "Region", "Total")
    check("% of grand: cells are shares (North A 18.52) and corner Total is 100",
          n is not None and abs(n["A"] - 18.52) < 0.01
          and t is not None and abs(t["Total"] - 100) < 0.01, grid.to_string())
    check("% note names the scope", "% of the grand total" in text(j), text(j)[:300])
else:
    check("% of grand ran", False, str(j)[:200])

j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                       "pivot_column": "Product", "value_column": "Price",
                       "percent_of": "row"}])
if grid is not None:
    n = row(grid, "Region", "North")
    check("% of row: North A 66.67 and every row's Total is 100",
          n is not None and abs(n["A"] - 66.67) < 0.01
          and (grid["Total"].dropna().sub(100).abs() < 0.01).all(), grid.to_string())
else:
    check("% of row ran", False, str(j)[:200])

# ---- (e) date bucketing (explicit month/quarter + auto-month) --------------------------
j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Date"],
                       "value_column": "Price", "date_bucket": "month"}])
if grid is not None:
    col = "Date (month)"
    check("month bucket: mixed-format dates land in the right months",
          col in grid.columns
          and row(grid, col, "2026-01")["sum_of_Price"] == 150
          and row(grid, col, "2026-02")["sum_of_Price"] == 200
          and row(grid, col, "2026-04")["sum_of_Price"] == 40, grid.to_string())
    check("month bucket: blank date -> '(blank)', garbage -> '(not a date)' + noted",
          row(grid, col, "(blank)") is not None
          and row(grid, col, "(not a date)") is not None
          and "recognizable date" in text(j), text(j)[:300])
else:
    check("month bucket ran", False, str(j)[:200])

j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Date"],
                       "value_column": "Price", "date_bucket": "quarter"}])
if grid is not None:
    col = "Date (quarter)"
    check("quarter bucket: Q1 350 / Q2 40",
          row(grid, col, "2026-Q1")["sum_of_Price"] == 350
          and row(grid, col, "2026-Q2")["sum_of_Price"] == 40, grid.to_string())
else:
    check("quarter bucket ran", False, str(j)[:200])

j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Date"],
                       "value_column": "Price"}])
check("date dim without a bucket auto-groups by month and SAYS so",
      grid is not None and "Date (month)" in grid.columns
      and "grouped 'Date' by month" in text(j),
      (grid.to_string() if grid is not None else str(j)[:200]))

# ---- (f) count without a value column + show_totals false ------------------------------
j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                       "agg_func": "count"}])
check("count with no value column = rows per group (North 3, South 4, Total 7)",
      grid is not None and row(grid, "Region", "North")["count_of_rows"] == 3
      and row(grid, "Region", "South")["count_of_rows"] == 4
      and row(grid, "Region", "Total")["count_of_rows"] == 7,
      (grid.to_string() if grid is not None else str(j)[:200]))

j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                       "value_column": "Price", "show_totals": False}])
check("show_totals false -> no Total row",
      grid is not None and row(grid, "Region", "Total") is None,
      (grid.to_string() if grid is not None else str(j)[:200]))

# ---- (g) live mode: GROUPBY / PIVOTBY formula + source sheet + M365 warning ------------
j, book, _ = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                       "value_column": "Price", "live": True}])
ok = j.get("status") == "ok" and book is not None
check("live 1-D: runs", ok, str(j)[:200])
if ok:
    main_ws = book[book.sheetnames[0]]
    f = main_ws["A1"].value
    check("live 1-D: A1 holds =GROUPBY(...) and the grid cells were cleared",
          isinstance(f, str) and f.startswith("=GROUPBY(")
          and main_ws["A2"].value is None and main_ws["B1"].value is None, repr(f))
    check("live 1-D: source data sheet exists with the raw rows",
          any("Pivot Data" in s for s in book.sheetnames)
          and book[[s for s in book.sheetnames if "Pivot Data" in s][0]].max_row == 8,
          str(book.sheetnames))
    check("live 1-D: response warns about Microsoft 365",
          "Microsoft 365" in text(j), text(j)[:300])
    check("live 1-D: preview still shows COMPUTED values",
          any(r.get("sum_of_Price") == 150 for tbl in (j.get("preview") or [])
              for r in tbl.get("sample_rows", [])), str(j.get("preview"))[:300])

j, book, _ = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                       "pivot_column": "Product", "value_column": "Price", "live": True}])
check("live 2-D: A1 holds =PIVOTBY(...)",
      book is not None and isinstance(book[book.sheetnames[0]]["A1"].value, str)
      and book[book.sheetnames[0]]["A1"].value.startswith("=PIVOTBY("),
      str(j)[:200])

# ---- (h) standard workbook invariant: messy 200-row Sales, Total == raw sum ------------
if WB_PATH.exists():
    raw_sales = pd.read_excel(WB_PATH, sheet_name="Sales")
    truth = pd.to_numeric(raw_sales["Qty"].astype(object), errors="coerce").sum()
    j, _, grid = run_ops([{"action": "pivot_summary", "group_by": ["Region"],
                           "value_column": "Qty"}],
                         data=WB_PATH.read_bytes(), fname="standard_test_workbook.xlsx")
    ok = grid is not None
    check("standard wb: pivot of 200 messy rows runs", ok, str(j)[:200])
    if ok:
        t = row(grid, "Region", "Total")
        check("standard wb: Total row equals the raw numeric sum of Qty",
              t is not None and abs(float(t["sum_of_Qty"]) - float(truth)) < 0.01,
              f"grid {t} vs truth {truth}")
        check("standard wb: blank regions grouped as '(blank)'",
              row(grid, "Region", "(blank)") is not None, grid["Region"].to_string())
else:
    check("standard workbook present", False, str(WB_PATH))

# ---- (i) plan validator knows pivot_column --------------------------------------------
tables = {"t": pd.DataFrame({"Region": ["N"], "Price": [1]})}
miss = _missing_columns([{"action": "pivot_summary", "group_by": ["Region"],
                          "pivot_column": "Zone", "value_column": "Price"}], tables)
check("validator flags a phantom pivot_column pre-flight", miss == ["Zone"], str(miss))

# ---- (j) failure / honesty paths (direct calls: exact messages) ------------------------
DF = pd.DataFrame({"Region": ["N", "S"], "Product": ["A", "B"], "Price": [1, 2]})


def err(op: dict, df: pd.DataFrame = DF) -> str:
    try:
        pivot_summary(df, op)
        return ""
    except OperationError as e:
        return str(e)


check("no group_by -> friendly ask",
      "at least one field" in err({"value_column": "Price"}))
check("value == group field -> refused",
      "both a grouping field" in err({"group_by": ["Price"], "value_column": "Price"}))
check("sum of a text column -> honest 'looks like text'",
      "looks like text" in err({"group_by": ["Region"], "value_column": "Product"}))
check("unknown agg -> friendly list",
      "sum, average, count" in err({"group_by": ["Region"], "value_column": "Price",
                                    "agg_func": "median"}))
check("% of an average -> refused as not meaningful",
      "isn't meaningful" in err({"group_by": ["Region"], "value_column": "Price",
                                 "agg_func": "average", "percent_of": "grand"}))
check("live + percent -> honest decline",
      "isn't available as a live" in err({"group_by": ["Region"], "value_column": "Price",
                                          "percent_of": "grand", "live": True}))
check("date_bucket on a non-date field -> honest error",
      "look like dates" in err({"group_by": ["Region"], "value_column": "Price",
                                "date_bucket": "month"}))
check("empty table -> friendly error",
      "no data rows" in err({"group_by": ["Region"], "value_column": "Price"},
                            pd.DataFrame(columns=["Region", "Price"])))
check("missing value column for sum -> asks which column",
      "Which column" in err({"group_by": ["Region"]}))

j, _, _ = run_ops([{"action": "pivot_summary", "group_by": ["Zone"],
                    "value_column": "Price"}])
check("HTTP: phantom column -> friendly failure, not a 500",
      j.get("status") != "ok"
      and "Zone" in (j.get("error") or "") + (j.get("clarification") or ""), str(j)[:300])

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
