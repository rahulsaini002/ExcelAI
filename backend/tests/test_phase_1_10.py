"""ENGINE PHASE 1.10 — fill series, legacy formats, named ranges (NO AI).

Series correctness (numbers-as-column, months-as-sheet, every-Monday dates), .xls and
.ods loading through /inspect, named ranges (defined name in the saved workbook +
same-plan {Alias:} use in a Formula-Generator op), the chained-plan validator fix,
and the failure paths.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_1_10.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p110.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import llm  # noqa: E402
from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"
NORMAL = (TESTS / "files" / "normal.xlsx").read_bytes()  # Date, Item, Qty, Price × 3 rows


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def run_ops(ops: list[dict], data: bytes = NORMAL, fname: str = "normal.xlsx"):
    sid = f"p110-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", (fname, data, OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": ops})})
    j = r.json()
    book = None
    if j.get("status") == "ok" and j.get("download_id"):
        book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
    return j, book


def sample(j: dict, column: str, sheet_needle: str = "") -> list:
    for t in j.get("preview") or []:
        if sheet_needle in t.get("name", ""):
            return [row.get(column) for row in t.get("sample_rows", [])]
    return []


print("ENGINE PHASE 1.10 — series / legacy formats / named ranges (no AI)\n")

# ---- (a) fill series -----------------------------------------------------------------------
j, _ = run_ops([{"action": "fill_series", "series_type": "numbers", "name": "No."}])
check("row-numbering column (fits table)", j.get("status") == "ok"
      and sample(j, "No.")[:3] == [1, 2, 3], f"{sample(j, 'No.')} {j.get('error', '')[:100]}")

j, book = run_ops([{"action": "fill_series", "series_type": "months", "name": "Month"}])
check("12 months don't fit 3 rows -> new sheet", j.get("status") == "ok"
      and "new sheet" in (j.get("explanation") or ""), (j.get("explanation") or j.get("error", ""))[:160])
if book:
    m = [s for s in book.sheetnames if "Month" in s]
    check("Month sheet holds January..December", bool(m)
          and book[m[0]].cell(row=2, column=1).value == "January"
          and book[m[0]].max_row == 13, str(book.sheetnames))

j, book = run_ops([{"action": "fill_series", "series_type": "dates", "every": "monday",
                    "start_date": "2026-07-01", "count": 5, "name": "Mondays"}])
check("every-Monday series executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    m = [s for s in book.sheetnames if "Mondays" in s]
    if m:
        vals = [book[m[0]].cell(row=r, column=1).value for r in range(2, 7)]
        ok = all(pd.Timestamp(v).dayofweek == 0 for v in vals) and pd.Timestamp(vals[0]) >= pd.Timestamp("2026-07-01")
        deltas = [(pd.Timestamp(vals[i + 1]) - pd.Timestamp(vals[i])).days for i in range(4)]
        check("all five are consecutive Mondays", ok and deltas == [7, 7, 7, 7],
              f"vals={vals} deltas={deltas}")

j, _ = run_ops([{"action": "fill_series", "series_type": "numbers", "step": 0}])
check("zero step -> clean error", j.get("status") != "ok" and "non-zero" in json.dumps(j),
      json.dumps(j)[:160])

j, _ = run_ops([{"action": "fill_series", "series_type": "fibonacci"}])
check("unknown series type -> clean error", j.get("status") != "ok" and "months" in json.dumps(j),
      json.dumps(j)[:160])

# ---- (b) legacy formats through /inspect ---------------------------------------------------
for ext in ("xls", "ods"):
    p = TESTS / "files" / f"legacy.{ext}"
    if not p.exists():
        print(f"  SKIP  .{ext} load (fixture missing)")
        continue
    r = client.post("/inspect", files=[("files", (p.name, p.read_bytes(), OCT))])
    jj = r.json()
    t = jj["tables"][0] if r.status_code == 200 and jj.get("tables") else {}
    check(f".{ext} loads (3 rows, Region/Amount headers)",
          r.status_code == 200 and t.get("row_count") == 3
          and [c["name"] for c in t.get("columns", [])] == ["Region", "Amount"],
          r.text[:160])

# ---- (c) named ranges -----------------------------------------------------------------------
j, book = run_ops([
    {"action": "name_range", "range_name": "Prices", "column": "Price"},
    {"action": "add_formula_column", "name": "Share", "formula": "{Price} / SUM({Prices:})"},
])
check("name_range + formula using {Prices:} in the SAME plan", j.get("status") == "ok",
      j.get("error", "")[:180])
if j.get("status") == "ok":
    vals = [v for v in sample(j, "Share") if v is not None]
    check("Share values sum to 1 (alias resolved to the Price column)",
          vals and abs(sum(vals) - 1.0) < 1e-9, f"vals={vals}")
if book:
    dn = book.defined_names.get("Prices") if hasattr(book.defined_names, "get") else None
    if dn is None and "Prices" in book.defined_names:
        dn = book.defined_names["Prices"]
    check("saved workbook carries the defined name over $D$2:$D$4",
          dn is not None and "$D$2:$D$4" in str(dn.attr_text), str(dn.attr_text if dn else None))

j, book = run_ops([
    {"action": "name_range", "range_name": "Prices", "column": "Price"},
    {"action": "layout_format", "title": "Priced"},
])
if book:
    dn = book.defined_names["Prices"] if "Prices" in book.defined_names else None
    check("defined name SHIFTED under a title ($D$3:$D$5)",
          dn is not None and "$D$3:$D$5" in str(dn.attr_text), str(dn.attr_text if dn else None))

j, _ = run_ops([{"action": "name_range", "range_name": "My Prices", "column": "Price"}])
check("space in range name -> clean error", j.get("status") != "ok"
      and "underscores" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_ops([{"action": "name_range", "range_name": "Qty", "column": "Price"}])
check("range name colliding with a column -> clean error", j.get("status") != "ok"
      and "already a column" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_ops([{"action": "name_range", "range_name": "Prices", "column": "Salary"}])
check("naming a missing column -> clean error", j.get("status") != "ok"
      and "Salary" in json.dumps(j), json.dumps(j)[:160])

# ---- validator fix: plans that CREATE then USE a name ---------------------------------------
real = llm.parse_instruction
try:
    llm.parse_instruction = lambda *a, **k: {
        "operations": [
            {"action": "add_formula_column", "name": "Total", "formula": "{Qty} * {Price}"},
            {"action": "sort", "columns": ["Total"], "orders": ["desc"]},
        ], "confidence": 90}
    r = client.post("/process", data={"instruction": "total then sort",
                                      "session_id": f"p110-{uuid.uuid4().hex[:8]}"},
                    files=[("files", ("normal.xlsx", NORMAL, OCT))])
    j = r.json()
    check("chained plan (create column -> sort by it) passes the validator",
          j.get("status") == "ok", json.dumps(j)[:180])

    llm.parse_instruction = lambda *a, **k: {
        "operations": [
            {"action": "name_range", "range_name": "Prices", "column": "Price"},
            {"action": "add_formula_column", "name": "Share", "formula": "{Price} / SUM({Prices:})"},
        ], "confidence": 90}
    r = client.post("/process", data={"instruction": "name and share",
                                      "session_id": f"p110-{uuid.uuid4().hex[:8]}"},
                    files=[("files", ("normal.xlsx", NORMAL, OCT))])
    j = r.json()
    check("named-range alias passes the validator through /process",
          j.get("status") == "ok", json.dumps(j)[:180])
finally:
    llm.parse_instruction = real

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
