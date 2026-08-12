"""ENGINE PHASE 1.2 — conditional formatting, Hands-layer verification (NO AI).

Runs hand-built conditional_format plans through the real API (/inspect -> /execute ->
/download) and reads the conditional-formatting RULES back out of the saved workbook
with openpyxl — checking rule type, operator, target range, and the honest match-count
note — plus every failure path (unknown rule, bad color, missing bound, bad column).

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_1_2.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p12.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"
WB = (TESTS / "standard_test_workbook.xlsx").read_bytes()


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def run_cf(op: dict):
    """inspect -> execute one conditional_format op -> (response json, Sales worksheet)."""
    sid = f"p12-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", ("standard_test_workbook.xlsx", WB, OCT))])
    assert r.status_code == 200, r.text[:200]
    plan = json.dumps({"operations": [{"action": "conditional_format", **op}]})
    r = client.post("/execute", data={"session_id": sid, "plan": plan})
    j = r.json()
    ws = None
    if j.get("status") == "ok" and j.get("download_id"):
        book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
        name = next((s for s in book.sheetnames if "Sales" in s), book.sheetnames[0])
        ws = book[name]
    return j, ws


def cf_rules(ws) -> list[tuple[str, str, str]]:
    """All CF rules in the sheet as (range, rule_type, operator-or-'')."""
    out = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            out.append((str(cf.sqref), rule.type or "", getattr(rule, "operator", "") or ""))
    return out


print("ENGINE PHASE 1.2 — conditional formatting (Hands layer, no AI)\n")

# ---- happy: one of each rule family ------------------------------------------------------
j, ws = run_cf({"columns": ["Price"], "rule_type": "greater_than", "value": 3000, "color": "green"})
check("greater_than executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    rules = cf_rules(ws)
    check("greater_than -> cellIs/greaterThan on E2:E201",
          ("E2:E201", "cellIs", "greaterThan") in rules, str(rules))
    check("note reports live match count", "match" in (j.get("explanation") or "")
          and "stays live" in (j.get("explanation") or ""), j.get("explanation", "")[:160])

j, ws = run_cf({"columns": ["Product"], "rule_type": "text_contains", "value": "widget", "color": "orange"})
check("text_contains executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    check("text_contains -> expression rule on C2:C201",
          any(r[0] == "C2:C201" and r[1] == "expression" for r in cf_rules(ws)), str(cf_rules(ws)))

j, ws = run_cf({"columns": ["Product"], "rule_type": "duplicates", "color": "red"})
check("duplicates executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    check("duplicates -> duplicateValues rule",
          any(r[1] == "duplicateValues" for r in cf_rules(ws)), str(cf_rules(ws)))

j, ws = run_cf({"columns": ["Region"], "rule_type": "unique"})
check("unique executes -> uniqueValues rule",
      j.get("status") == "ok" and ws is not None and any(r[1] == "uniqueValues" for r in cf_rules(ws)),
      j.get("error", "")[:140])

j, ws = run_cf({"columns": ["Price"], "rule_type": "color_scale"})
check("color_scale executes -> colorScale rule",
      j.get("status") == "ok" and ws is not None and any(r[1] == "colorScale" for r in cf_rules(ws)),
      j.get("error", "")[:140])

j, ws = run_cf({"columns": ["Qty"], "rule_type": "data_bars", "color": "blue"})
check("data_bars executes -> dataBar rule",
      j.get("status") == "ok" and ws is not None and any(r[1] == "dataBar" for r in cf_rules(ws)),
      j.get("error", "")[:140])

j, ws = run_cf({"columns": ["Qty"], "rule_type": "icon_set", "icons": 3})
check("icon_set executes -> iconSet rule",
      j.get("status") == "ok" and ws is not None and any(r[1] == "iconSet" for r in cf_rules(ws)),
      j.get("error", "")[:140])

j, ws = run_cf({"columns": ["Price"], "rule_type": "top_n", "count": 10, "color": "green"})
check("top_n executes -> top10 rule", j.get("status") == "ok" and ws is not None
      and any(r[1] == "top10" for r in cf_rules(ws)), j.get("error", "")[:140])

j, ws = run_cf({"columns": ["Qty"], "rule_type": "blanks"})
check("blanks executes -> containsBlanks rule",
      j.get("status") == "ok" and ws is not None and any(r[1] == "containsBlanks" for r in cf_rules(ws)),
      j.get("error", "")[:140])
if j.get("status") == "ok":
    # honest count: recompute blanks straight from the workbook
    df = pd.read_excel(io.BytesIO(WB), sheet_name="Sales")
    blanks = int(df["Qty"].isna().sum())
    check("blank-count note matches pandas truth",
          f"{blanks:,} match" in (j.get("explanation") or ""), j.get("explanation", "")[:160])

j, ws = run_cf({"columns": ["Price"], "rule_type": "between", "value": 1000, "value2": 2000, "color": "yellow"})
check("between executes -> cellIs/between", j.get("status") == "ok" and ws is not None
      and any(r[1] == "cellIs" and r[2] == "between" for r in cf_rules(ws)), j.get("error", "")[:140])

j, ws = run_cf({"columns": ["Date"], "rule_type": "date_after", "value": "2026-03-01", "color": "blue"})
check("date_after executes -> expression rule", j.get("status") == "ok" and ws is not None
      and any(r[1] == "expression" for r in cf_rules(ws)), j.get("error", "")[:140])

j, ws = run_cf({"columns": ["Price"], "rule_type": "formula", "formula": "{Price} > 100 * {Qty}", "color": "red"})
check("formula rule executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    got = [r for r in cf_rules(ws) if r[1] == "expression"]
    check("formula rule -> expression with translated refs", bool(got), str(cf_rules(ws)))

# ---- edge --------------------------------------------------------------------------------
j, ws = run_cf({"columns": ["Qty", "Price"], "rule_type": "greater_than", "value": 10, "color": "purple"})
check("multi-column rule lands on BOTH ranges", j.get("status") == "ok" and ws is not None
      and any(r[0] == "D2:D201" for r in cf_rules(ws)) and any(r[0] == "E2:E201" for r in cf_rules(ws)),
      str(cf_rules(ws) if ws else j)[:160])

j, ws = run_cf({"columns": ["Price"], "rule_type": "top_n", "count": 10, "percent": True})
check("top 10 PERCENT variant executes", j.get("status") == "ok"
      and "10" in (j.get("explanation") or ""), j.get("error", "")[:140])

# ---- failure / honesty -------------------------------------------------------------------
j, _ = run_cf({"columns": ["Price"], "rule_type": "sparkle_rainbow"})
check("unknown rule -> clean error listing what's possible",
      j.get("status") != "ok" and "I can do" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_cf({"columns": ["Price"], "rule_type": "greater_than", "value": 100, "color": "vermilion"})
check("unknown color -> clean error naming colors",
      j.get("status") != "ok" and "green" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_cf({"columns": ["Price"], "rule_type": "between", "value": 100})
check("between without second bound -> clean error",
      j.get("status") != "ok" and "both bounds" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_cf({"columns": ["Price"], "rule_type": "greater_than"})
check("comparison without value -> clean error",
      j.get("status") != "ok" and "needs a value" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_cf({"columns": ["Salary"], "rule_type": "greater_than", "value": 100})
check("missing column -> clean error naming it",
      j.get("status") != "ok" and "Salary" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_cf({"columns": ["Qty"], "rule_type": "icon_set", "icons": 7})
check("7-icon set -> clean error (3/4/5)",
      j.get("status") != "ok" and "3, 4, or 5" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_cf({"columns": ["Date"], "rule_type": "date_before", "value": "not-a-date"})
check("unreadable date -> clean error",
      j.get("status") != "ok" and "as a date" in json.dumps(j), json.dumps(j)[:160])

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
