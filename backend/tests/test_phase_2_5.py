"""ENGINE PHASE 2.5 — dashboards (NO AI).

Verify-and-harden the one-page dashboard: KPIs computed correctly (sum/mean/count/
count_distinct/min/max × number/currency/percent), the FULL chart family usable inside
a dashboard (2.5 now reuses the 2.4 chart validation), a written summary whose figures
are TRUSTED (built from the computed KPIs, not the model — so it can't fabricate a
total), non-overlapping layout on a dedicated Dashboard sheet, clean regeneration when
the data changes, the aggregate→dashboard sequence, and the honest failure paths.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_2_5.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
import zipfile
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p25.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import app.main as m  # noqa: E402
from app.db import init_db  # noqa: E402
from app.executor import execute_multi  # noqa: E402
from app.operations.base import OperationError  # noqa: E402
from app.operations.dashboard import dashboard  # noqa: E402

init_db()
client = TestClient(m.app)
passed = failed = 0
OCT = "application/octet-stream"

DF = pd.DataFrame({
    "Month": ["Jan", "Feb", "Mar", "Apr"],
    "Revenue": [100, 200, 150, 250],
    "Orders": [3, 5, 4, 6],
    "MonthNum": [1, 2, 3, 4],
})


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def err(op: dict, df: pd.DataFrame = DF) -> str:
    try:
        dashboard(df, op)
        return ""
    except OperationError as e:
        return str(e)


print("ENGINE PHASE 2.5 — dashboards (no AI)\n")

# ---- (a) KPIs compute correctly across aggs + formats ----
note, d = dashboard(DF, {"dashboard_title": "Shop", "kpis": [
    {"label": "Total Rev", "agg": "sum", "column": "Revenue", "format": "currency"},
    {"label": "Avg Orders", "agg": "mean", "column": "Orders"},
    {"label": "Rows", "agg": "count"},
    {"label": "Distinct Months", "agg": "count_distinct", "column": "Month"},
    {"label": "Min Rev", "agg": "min", "column": "Revenue"},
    {"label": "Max Rev", "agg": "max", "column": "Revenue"},
]})
kpis = {k["label"]: k["value"] for k in d["kpis"]}
check("KPI sum as currency (₹700)", kpis["Total Rev"] == "₹700", str(kpis))
check("KPI mean (4.5)", kpis["Avg Orders"] == "4.50" or kpis["Avg Orders"] == "4.5", str(kpis))
check("KPI plain row count (4)", kpis["Rows"] == "4", str(kpis))
check("KPI count_distinct (4)", kpis["Distinct Months"] == "4", str(kpis))
check("KPI min/max (100 / 250)", kpis["Min Rev"] == "100" and kpis["Max Rev"] == "250", str(kpis))

_, dp = dashboard(pd.DataFrame({"Rate": [0.25], "X": [1]}),
                  {"kpis": [{"label": "R", "agg": "sum", "column": "Rate", "format": "percent"}]})
check("KPI percent format", dp["kpis"][0]["value"] == "0.2%" or dp["kpis"][0]["value"].endswith("%"), str(dp["kpis"]))

# uncomputable KPI (sum of text) -> '—' and the note says so
_, du = dashboard(DF, {"kpis": [{"label": "Bad", "agg": "sum", "column": "Month"}]})
check("uncomputable KPI shows '—'", du["kpis"][0]["value"] == "—", str(du["kpis"]))

# ---- (b) trusted, ACCURATE written summary (can't fabricate) ----
_, ds = dashboard(DF, {
    "kpis": [{"label": "Total Rev", "agg": "sum", "column": "Revenue", "format": "currency"}],
    "summary": "Revenue is fabricated as ₹99999 by the model",  # a lying narrative
})
check("summary carries the REAL figure from KPIs (₹700)", "₹700" in ds["summary"], ds["summary"])
check("summary states the real row count", "4 rows" in ds["summary"], ds["summary"])
check("summary keeps the qualitative narrative alongside facts",
      "fabricated" in ds["summary"] and "By the numbers" in ds["summary"], ds["summary"])
# the trusted figure is present even though the model's number was different -> not faked
check("trusted figures don't echo the model's wrong number as a KPI",
      ds["kpis"][0]["value"] == "₹700", str(ds["kpis"]))

# ---- (c) the FULL chart family works inside a dashboard ----
for ct, xcol in [("radar", "Month"), ("doughnut", "Month"), ("stock", "Month"),
                 ("scatter", "MonthNum"), ("bar", "Month"), ("line", "Month")]:
    yc = ["Revenue", "Orders", "MonthNum"] if ct == "stock" else ["Revenue"]
    _, dc = dashboard(DF, {"charts": [{"chart_type": ct, "x_column": xcol, "y_columns": yc}]})
    check(f"dashboard chart '{ct}' validated", dc["charts"][0]["chart_type"] == ct, str(dc["charts"]))

# bubble inside a dashboard (needs numeric x + size)
_, db = dashboard(DF, {"charts": [{"chart_type": "bubble", "x_column": "MonthNum",
                                   "y_columns": ["Revenue"], "size_column": "Orders"}]})
check("dashboard bubble chart with size_column", db["charts"][0]["size_column"] == "Orders", str(db["charts"]))

# unsupported chart in a dashboard -> honest fallback (never silent)
check("dashboard unsupported chart -> honest fallback",
      "treemap" in err({"charts": [{"chart_type": "treemap", "x_column": "Month", "y_columns": ["Revenue"]}]}))
# scatter with a TEXT x-axis in a dashboard is caught
check("dashboard scatter needs numeric x",
      "must be numeric" in err({"charts": [{"chart_type": "scatter", "x_column": "Month", "y_columns": ["Revenue"]}]}))

# ---- (d) data unchanged + directive flows through executor ----
res, _, notes, render = execute_multi({"t": DF.copy()}, "t", [{
    "action": "dashboard", "kpis": [{"label": "R", "agg": "sum", "column": "Revenue"}],
    "charts": [{"chart_type": "doughnut", "x_column": "Month", "y_columns": ["Revenue"]}]}])
check("dashboard leaves data unchanged", list(res.columns) == list(DF.columns) and len(res) == len(DF))
check("dashboard directive emitted", any(r.get("type") == "dashboard" for r in render))

# ---- (e) aggregate -> dashboard: KPIs/charts reflect the aggregated data ----
many = pd.DataFrame({"Region": ["N", "N", "S", "S"], "Sales": [10, 20, 30, 40]})
_, _, _, render2 = execute_multi({"t": many}, "t", [
    {"action": "aggregate", "agg_func": "sum", "agg_column": "Sales", "group_by": ["Region"]},
    {"action": "dashboard", "kpis": [{"label": "Total", "agg": "sum", "column": "sum_of_Sales"}],
     "charts": [{"chart_type": "bar", "x_column": "Region", "y_columns": ["sum_of_Sales"]}]},
])
ddir = next(r for r in render2 if r.get("type") == "dashboard")
check("aggregate->dashboard: KPI sums the 2 grouped rows (100)",
      ddir["kpis"][0]["value"] == "100", str(ddir["kpis"]))
check("aggregate->dashboard: chart references the aggregated column",
      ddir["charts"][0]["y_columns"] == ["sum_of_Sales"], str(ddir["charts"]))

# ---- (f) failure paths ----
check("empty data declined", "no data" in err({"kpis": [{"label": "R", "agg": "sum", "column": "Revenue"}]}, DF.iloc[0:0]))
check("bad KPI column named", "Nope" in err({"kpis": [{"label": "x", "agg": "sum", "column": "Nope"}]}))
check("empty dashboard (no kpis/charts) declined", "at least one" in err({}))

# ---- (g) HTTP round-trip: Dashboard sheet + chart parts, layout non-overlapping ----
buf = io.BytesIO()
with pd.ExcelWriter(buf, engine="openpyxl") as w:
    DF.to_excel(w, index=False, sheet_name="Sheet1")
sid = f"p25-{uuid.uuid4().hex[:10]}"
client.post("/inspect", data={"session_id": sid}, files=[("files", ("d.xlsx", buf.getvalue(), OCT))])
plan = {"operations": [{"action": "dashboard", "dashboard_title": "Overview",
        "kpis": [{"label": "Total Rev", "agg": "sum", "column": "Revenue", "format": "currency"}],
        "charts": [{"chart_type": "radar", "x_column": "Month", "y_columns": ["Revenue"]},
                   {"chart_type": "doughnut", "x_column": "Month", "y_columns": ["Orders"]}],
        "summary": "Solid quarter."}]}
r = client.post("/execute", data={"session_id": sid, "plan": json.dumps(plan)})
j = r.json()
book = None
if j.get("status") == "ok" and j.get("download_id"):
    raw = client.get(f"/download/{j['download_id']}").content
    book = load_workbook(io.BytesIO(raw))
    parts = zipfile.ZipFile(io.BytesIO(raw)).namelist()
check("HTTP: status ok", j.get("status") == "ok", str(j)[:150])
check("HTTP: Dashboard sheet + data sheet both present",
      book is not None and "Dashboard" in book.sheetnames and "Sheet1" in book.sheetnames,
      str(book.sheetnames) if book else "no book")
if book is not None:
    dash = book["Dashboard"]
    check("HTTP: title in A1", dash["A1"].value == "Overview", str(dash["A1"].value))
    check("HTTP: two charts embedded, non-overlapping",
          len(dash._charts) == 2 and len({str(c.anchor) for c in dash._charts}) == 2, str(len(dash._charts)))
    check("HTTP: .xlsx contains chart parts", any("chart" in p for p in parts), "")
    check("HTTP: original data sheet intact", book["Sheet1"]["A1"].value == "Month", str(book["Sheet1"]["A1"].value))

# regenerates cleanly (one Dashboard sheet, updated figures) on changed data
_, d2 = dashboard(pd.DataFrame({"Month": ["Jan"], "Revenue": [9000]}),
                  {"kpis": [{"label": "Total Rev", "agg": "sum", "column": "Revenue", "format": "currency"}]})
out2, _, _ = m._serialize(pd.DataFrame({"Month": ["Jan"], "Revenue": [9000]}), "x.csv", "xlsx", [d2])
wb2 = load_workbook(io.BytesIO(out2))
check("regen: exactly one Dashboard sheet", wb2.sheetnames.count("Dashboard") == 1, str(wb2.sheetnames))
check("regen: KPI reflects changed data (₹9.0K)", d2["kpis"][0]["value"] == "₹9.0K", str(d2["kpis"]))

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
