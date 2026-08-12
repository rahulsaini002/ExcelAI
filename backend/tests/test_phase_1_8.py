"""ENGINE PHASE 1.8 — Goal Seek, Hands-layer verification (NO AI).

Hand-built goal_seek plans through the real API: exact linear solves, nonlinear
bisection, the honest no-solution decline, the single-variable rule, the 'Goal Seek'
working sheet in the saved file, data immutability, and the {var} validator exemption.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_1_8.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p18.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import llm  # noqa: E402
from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def volumes_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Qty"])
    for p, q in [("A", 40), ("B", 100), ("C", 60)]:  # SUM(Qty) = 200
        ws.append([p, q])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_gs(op: dict):
    sid = f"p18-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", ("volumes.xlsx", volumes_xlsx(), OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid,
                                      "plan": json.dumps({"operations": [{"action": "goal_seek", **op}]})})
    j = r.json()
    book = None
    if j.get("status") == "ok" and j.get("download_id"):
        book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
    return j, book


print("ENGINE PHASE 1.8 — Goal Seek (Hands layer, no AI)\n")

# ---- linear solve: price x SUM(Qty)=200 -> target 1,000,000 => price 5,000 ---------------
j, book = run_gs({"formula": "{var} * SUM({Qty:})", "target": 1_000_000,
                  "variable_name": "price"})
check("linear goal executes", j.get("status") == "ok", j.get("error", "")[:180])
note = j.get("explanation") or ""
check("found value is exact (5,000)", "5,000.0000" in note, note[:160])
check("note shows target + achieved + method", "1,000,000.00" in note and "secant" in note, note[:200])
if book:
    gs = [s for s in book.sheetnames if "Goal Seek" in s]
    check("'Goal Seek' working sheet saved", bool(gs), str(book.sheetnames))
    if gs:
        w = book[gs[0]]
        rows = {w.cell(row=r, column=1).value: w.cell(row=r, column=2).value
                for r in range(1, 8)}
        check("working sheet has formula/target/found/achieved",
              rows.get("Formula") == "{var} * SUM({Qty:})" and rows.get("Target") == 1_000_000
              and abs((rows.get("Found value") or 0) - 5000) < 1e-6,
              str(rows))
    data = book[next(s for s in book.sheetnames if "Goal Seek" not in s)]
    check("data unchanged (Qty intact)", data["B2"].value == 40 and data.max_row == 4,
          f"B2={data['B2'].value} rows={data.max_row}")

# ---- nonlinear: x^2 * SUM(Qty) = 800 => x = 2 ---------------------------------------------
j, _ = run_gs({"formula": "POWER({var}, 2) * SUM({Qty:})", "target": 800,
               "variable_name": "factor"})
check("nonlinear goal solves (x=2)", j.get("status") == "ok"
      and "2.0000" in (j.get("explanation") or ""), (j.get("explanation") or j.get("error", ""))[:180])

# ---- no solution: x^2 * positive = negative target ----------------------------------------
j, _ = run_gs({"formula": "POWER({var}, 2) * SUM({Qty:})", "target": -500,
               "variable_name": "factor"})
check("unreachable target -> honest decline", j.get("status") != "ok"
      and "unreachable" in json.dumps(j), json.dumps(j)[:200])

# ---- single-variable rule ------------------------------------------------------------------
j, _ = run_gs({"formula": "{var1} * {var2} * SUM({Qty:})", "target": 100})
check("two unknowns -> declined (single variable v1)", j.get("status") != "ok"
      and "ONE unknown" in json.dumps(j), json.dumps(j)[:200])

j, _ = run_gs({"formula": "SUM({Qty:}) * 2", "target": 100})
check("no {var} in formula -> clean ask", j.get("status") != "ok"
      and "{var}" in json.dumps(j), json.dumps(j)[:200])

j, _ = run_gs({"formula": "{var} * SUM({Qty:})"})
check("missing target -> clean ask", j.get("status") != "ok"
      and "number" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_gs({"formula": "{var} * SUM({Product:})", "target": 100})
check("text column in the formula -> clean error", j.get("status") != "ok", json.dumps(j)[:200])

# ---- {var} exempt from the phantom-column validator (via patched /process) -----------------
real = llm.parse_instruction
try:
    llm.parse_instruction = lambda *a, **k: {
        "operations": [{"action": "goal_seek", "formula": "{var} * SUM({Qty:})",
                        "target": 1_000_000, "variable_name": "price"}], "confidence": 90}
    r = client.post("/process", data={"instruction": "what price gives 1000000",
                                      "session_id": f"p18-{uuid.uuid4().hex[:8]}"},
                    files=[("files", ("volumes.xlsx", volumes_xlsx(), OCT))])
    j = r.json()
    check("{var} NOT flagged by the plan validator (full /process path ok)",
          j.get("status") == "ok" and "5,000" in (j.get("explanation") or ""),
          json.dumps(j)[:200])
finally:
    llm.parse_instruction = real

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
