"""ENGINE PHASE 2.4 — charts, full set (NO AI).

Every openpyxl chart type Sumio supports (bar/line/area/pie/doughnut/radar/stock +
scatter/bubble) builds a REAL chart that embeds in a valid .xlsx and opens cleanly;
the category vs XY data models are wired correctly (scatter/bubble need a numeric
x-axis; bubble needs a size); single-series charts (pie/doughnut) collapse extra
series; and every UNSUPPORTED type (histogram/waterfall/funnel/treemap/sunburst/
sparkline/gauge/heatmap/map) is declined with an honest fallback — never silently
swapped. Charts never change the data.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_2_4.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
import zipfile
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p24.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import app.main as m  # noqa: E402
from app.db import init_db  # noqa: E402
from app.executor import execute_multi  # noqa: E402
from app.operations.base import OperationError  # noqa: E402
from app.operations.chart import chart, CATEGORY_TYPES, XY_TYPES  # noqa: E402

init_db()
client = TestClient(app := m.app)
passed = failed = 0
OCT = "application/octet-stream"

DF = pd.DataFrame({
    "Label": ["A", "B", "C", "D", "E"],
    "X": [1.0, 2.0, 3.0, 4.0, 5.0],
    "Y": [10, 25, 18, 30, 22],
    "Size": [5, 9, 3, 7, 4],
    "High": [12, 26, 20, 33, 24], "Low": [8, 20, 15, 27, 19], "Close": [11, 24, 18, 30, 22],
})


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def err(op: dict, df: pd.DataFrame = DF) -> str:
    try:
        chart(df, op)
        return ""
    except OperationError as e:
        return str(e)


def op_for(ct: str) -> dict:
    op = {"action": "chart", "chart_type": ct, "chart_title": ct}
    if ct in ("scatter", "bubble"):
        op.update(x_column="X", y_columns=["Y"])
        if ct == "bubble":
            op["size_column"] = "Size"
    elif ct == "stock":
        op.update(x_column="Label", y_columns=["High", "Low", "Close"])
    else:
        op.update(x_column="Label", y_columns=["Y"])
    return op


print("ENGINE PHASE 2.4 — charts, full set (no AI)\n")

# ---- (a) every supported type builds a real, embedded, openable chart ----
for ct in CATEGORY_TYPES + XY_TYPES:
    note, d = chart(DF, op_for(ct))
    out, fname, media = m._serialize(DF, "c.csv", "xlsx", [d])
    wb = load_workbook(io.BytesIO(out))
    part = any("chart" in n for n in zipfile.ZipFile(io.BytesIO(out)).namelist())
    check(f"{ct}: builds + embeds + opens cleanly",
          d["chart_type"] == ct and len(wb.active._charts) == 1 and part
          and wb.active.max_row == len(DF) + 1 and fname.endswith(".xlsx"),
          f"charts={len(wb.active._charts)} part={part}")

# ---- (b) XY charts need a numeric x-axis ----
check("scatter: non-numeric x-axis declined",
      "must be numeric" in err({"chart_type": "scatter", "x_column": "Label", "y_columns": ["Y"]}))
check("scatter: numeric x-axis accepted",
      chart(DF, {"chart_type": "scatter", "x_column": "X", "y_columns": ["Y"]})[1]["chart_type"] == "scatter")
check("category chart accepts a TEXT x-axis (bar of Y by Label)",
      chart(DF, {"chart_type": "bar", "x_column": "Label", "y_columns": ["Y"]})[1]["x_column"] == "Label")

# ---- (c) bubble needs a size; falls back to a 2nd value column ----
check("bubble: no size column declined",
      "bubble" in err({"chart_type": "bubble", "x_column": "X", "y_columns": ["Y"]}).lower())
note, d = chart(DF, {"chart_type": "bubble", "x_column": "X", "y_columns": ["Y", "Size"]})
check("bubble: 2nd value column used as size",
      d["size_column"] == "Size" and d["y_columns"] == ["Y"], str(d))
check("bubble: explicit size_column honored",
      chart(DF, {"chart_type": "bubble", "x_column": "X", "y_columns": ["Y"], "size_column": "Size"})[1]["size_column"] == "Size")

# ---- (d) single-series charts collapse extra series ----
_, dp = chart(DF, {"chart_type": "pie", "x_column": "Label", "y_columns": ["Y", "Size"]})
check("pie: keeps a single series", dp["y_columns"] == ["Y"], str(dp))
_, dd = chart(DF, {"chart_type": "doughnut", "x_column": "Label", "y_columns": ["Y", "Size"]})
check("doughnut: keeps a single series", dd["y_columns"] == ["Y"], str(dd))

# ---- (e) synonyms ----
check("synonym: 'column' -> bar",
      chart(DF, {"chart_type": "column", "x_column": "Label", "y_columns": ["Y"]})[1]["chart_type"] == "bar")
check("synonym: 'donut' -> doughnut",
      chart(DF, {"chart_type": "donut", "x_column": "Label", "y_columns": ["Y"]})[1]["chart_type"] == "doughnut")

# ---- (f) unsupported types declined with an honest fallback ----
for bad, needle in [("histogram", "bar"), ("heatmap", "colour"), ("waterfall", "bar"),
                    ("funnel", "bar"), ("treemap", "pie"), ("sparkline", "line"),
                    ("gauge", "KPI"), ("sunburst", "pie")]:
    msg = err({"chart_type": bad, "x_column": "Label", "y_columns": ["Y"]})
    check(f"unsupported '{bad}' -> honest fallback ({needle})",
          bad in msg and needle in msg, msg[:90])
check("unsupported 'map' -> honest no-equivalent",
      "isn't something an Excel chart" in err({"chart_type": "map", "x_column": "Label", "y_columns": ["Y"]}))
check("unknown chart word -> lists supported types",
      "I can do" in err({"chart_type": "zigzag", "x_column": "Label", "y_columns": ["Y"]}))

# ---- (g) generic failures ----
check("empty data declined", "no data" in err(op_for("bar"), DF.iloc[0:0]))
check("missing x column named", "X2" in err({"chart_type": "bar", "x_column": "X2", "y_columns": ["Y"]}) or
      "don't see" in err({"chart_type": "bar", "x_column": "X2", "y_columns": ["Y"]}))
check("non-numeric value column declined",
      "must be numbers" in err({"chart_type": "bar", "x_column": "X", "y_columns": ["Label"]}))
check("no value column declined", "value column" in err({"chart_type": "bar", "x_column": "Label"}))

# ---- (h) chart never changes the data; reflects CURRENT (post-aggregate) data ----
res, _, notes, render = execute_multi({"t": DF.copy()}, "t", [op_for("radar")])
check("chart op leaves data unchanged", list(res.columns) == list(DF.columns) and len(res) == len(DF))
check("chart op emits a render directive", any(r.get("type") == "chart" for r in render))

agg = pd.DataFrame({"Month": ["Jan", "Jan", "Feb"], "Rev": [100, 50, 200]})
_, _, _, render2 = execute_multi({"t": agg}, "t", [
    {"action": "aggregate", "agg_func": "sum", "agg_column": "Rev", "group_by": ["Month"]},
    {"action": "chart", "chart_type": "doughnut", "x_column": "Month", "y_columns": ["sum_of_Rev"]},
])
cdir = next(r for r in render2 if r.get("type") == "chart")
check("chart reflects aggregated data (2 categories)", cdir["rows"] == 2, str(cdir["rows"]))

# ---- (i) HTTP round-trip: scatter + bubble + stock through the real API ----
def api(op: dict):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        DF.to_excel(w, index=False, sheet_name="Sheet1")
    sid = f"p24-{uuid.uuid4().hex[:10]}"
    client.post("/inspect", data={"session_id": sid},
                files=[("files", ("chart.xlsx", buf.getvalue(), OCT))])
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": [op]})})
    j = r.json()
    charts = 0
    if j.get("status") == "ok" and j.get("download_id"):
        wb = load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
        charts = len(wb.active._charts)
    return j, charts

for ct in ("scatter", "bubble", "stock", "radar", "doughnut"):
    j, charts = api(op_for(ct))
    check(f"HTTP {ct}: ok + chart embedded in the download", j.get("status") == "ok" and charts == 1, str(j)[:120])

j, _ = api({"action": "chart", "chart_type": "treemap", "x_column": "Label", "y_columns": ["Y"]})
check("HTTP unsupported treemap: friendly decline, not a 500",
      j.get("status") != "ok" and "treemap" in json.dumps(j).lower(), str(j)[:120])

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
