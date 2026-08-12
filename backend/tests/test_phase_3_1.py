"""ENGINE PHASE 3.1 — multi-file merge + smart mapping + FUZZY matching (NO AI).

Merge already unified same-meaning columns (synonym groups + case/space/punct), kept
different columns separate, and flagged type conflicts. Phase 3.1 adds typo-tolerant
FUZZY matching — automatic, high-confidence only, and honestly reported so nothing is
silently mis-joined:
  * lookup: a key with a typo ("Jon Smith") finds the close source key ("John Smith"),
    written as a STATIC value (a fuzzy match can't be a live Excel formula) and flagged;
    genuinely different keys stay unmatched (no false match).
  * merge: near-identical column names ("Custmer_ID" ~ "Customer_ID") unify (likely
    typo), reported; distinct columns are still kept separate.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_3_1.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p31.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import app.main as m  # noqa: E402
from app.db import init_db  # noqa: E402
from app.executor import execute_multi  # noqa: E402

init_db()
client = TestClient(m.app)
passed = failed = 0
OCT = "application/octet-stream"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def run(ops, tables):
    res, name, notes, render = execute_multi(tables, next(iter(tables)), ops)
    return res, name, " ".join(notes), render


print("ENGINE PHASE 3.1 — merge + smart mapping + fuzzy matching (no AI)\n")

# ============ FUZZY LOOKUP ============
people = pd.DataFrame({"Name": ["John Smith", "Jon Smith", "Mary Jones", "Zzxqwv"]})
roster = pd.DataFrame({"FullName": ["John Smith", "Mary Jones"], "Dept": ["Eng", "Sales"]})
LK = {"action": "lookup", "key_column": "Name", "source_sheet": "roster",
      "source_key_column": "FullName", "return_column": "Dept", "new_column": "Dept"}
res, name, note, render = run([LK], {"people": people.copy(), "roster": roster.copy()})
vals = dict(zip(res["Name"], res["Dept"]))
check("fuzzy lookup: exact key matches (John Smith → Eng)", vals["John Smith"] == "Eng", str(vals))
check("fuzzy lookup: a TYPO key matches (Jon Smith → Eng)", vals["Jon Smith"] == "Eng", str(vals))
check("fuzzy lookup: a genuinely different key stays UNMATCHED (no false match)",
      vals["Zzxqwv"] == "Not found", str(vals))
check("fuzzy lookup: note flags the close-similarity match + example",
      "CLOSE SIMILARITY" in note and "Jon Smith" in note and "John Smith" in note, note)
check("fuzzy lookup: note says fuzzy rows are fixed values to verify",
      "fixed values" in note and "verify" in note, note)

# the saved file: the fuzzy row is a STATIC value, exact rows keep the LIVE formula
out, _, _ = m._serialize_workbook({name: res}, "x.xlsx", name, render)
ws = load_workbook(io.BytesIO(out)).active
dept_col = list(res.columns).index("Dept") + 1
r_exact = 2   # John Smith (row 1)
r_fuzzy = 3   # Jon Smith (row 2)
check("file: exact-match row keeps a live lookup formula",
      isinstance(ws.cell(row=r_exact, column=dept_col).value, str)
      and ws.cell(row=r_exact, column=dept_col).value.startswith("="), repr(ws.cell(row=r_exact, column=dept_col).value)[:40])
check("file: fuzzy-match row is a STATIC value (not a formula, not 'Not found')",
      ws.cell(row=r_fuzzy, column=dept_col).value == "Eng", repr(ws.cell(row=r_fuzzy, column=dept_col).value))

# high threshold: a loose resemblance is NOT matched (honesty — don't over-match)
p2 = pd.DataFrame({"Name": ["Cat"]})
s2 = pd.DataFrame({"FullName": ["Dog"], "V": ["x"]})
res2, _, note2, _ = run([{"action": "lookup", "key_column": "Name", "source_sheet": "s",
                          "source_key_column": "FullName", "return_column": "V", "new_column": "V"}],
                        {"p": p2.copy(), "s": s2.copy()})
check("fuzzy lookup: unrelated words are NOT force-matched (Cat ≠ Dog)",
      res2["V"].iloc[0] == "Not found" and "SIMILARITY" not in note2, str(res2["V"].tolist()))

# ============ FUZZY MERGE COLUMNS ============
A = pd.DataFrame({"Customer_ID": [1, 2], "Amount": [10, 20]})
B = pd.DataFrame({"Custmer_ID": [3, 4], "Amount": [30, 40]})   # typo in the ID column
res, name, note, _ = run([{"action": "merge", "merge_tables": ["A", "B"], "new_table": "all"}],
                         {"A": A.copy(), "B": B.copy()})
check("fuzzy merge: typo'd column names unify (one Customer_ID, not two)",
      "Customer_ID" in res.columns and "Custmer_ID" not in res.columns and len(res.columns) == 2, str(list(res.columns)))
check("fuzzy merge: all rows stacked (4)", len(res) == 4, str(len(res)))
check("fuzzy merge: note flags the likely-typo unification",
      "likely typos" in note and "Custmer_ID" in note, note)

# genuinely different columns are still kept SEPARATE (not fuzzy-merged)
C = pd.DataFrame({"Region": ["N"], "Sales": [5]})
D = pd.DataFrame({"Product": ["X"], "Sales": [7]})
res, _, note, _ = run([{"action": "merge", "merge_tables": ["C", "D"], "new_table": "cd"}],
                      {"C": C.copy(), "D": D.copy()})
check("distinct columns kept separate (Region & Product both present)",
      "Region" in res.columns and "Product" in res.columns, str(list(res.columns)))

# ============ EXISTING SMARTS still work (regression) ============
# synonym-group override + type-conflict flag
E = pd.DataFrame({"client_id": [1], "Val": [10]})
F = pd.DataFrame({"cust_no": [2], "Val": ["oops"]})   # Val is text here → type conflict
res, _, note, _ = run([{"action": "merge", "merge_tables": ["E", "F"], "new_table": "ef",
                        "column_groups": [{"name": "Customer_ID", "aliases": ["client_id", "cust_no"]}]}],
                      {"E": E.copy(), "F": F.copy()})
check("synonym override unifies client_id + cust_no → Customer_ID",
      "Customer_ID" in res.columns and "client_id" not in res.columns and "cust_no" not in res.columns,
      str(list(res.columns)))
check("type conflict on 'Val' is flagged", "numbers in some files and text in others" in note, note)

# ============ HTTP round-trip: fuzzy lookup end to end ============
def _xlsx(df):
    b = io.BytesIO()
    with pd.ExcelWriter(b, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Sheet1")
    return b.getvalue()

sid = f"p31-{uuid.uuid4().hex[:10]}"
client.post("/inspect", data={"session_id": sid}, files=[  # two separate FILES → table names = stems
    ("files", ("people.xlsx", _xlsx(people), OCT)),
    ("files", ("roster.xlsx", _xlsx(roster), OCT))])
r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": [
    {"action": "lookup", "key_column": "Name", "source_sheet": "roster",
     "source_key_column": "FullName", "return_column": "Dept", "new_column": "Dept", "table": "people"}]})})
j = r.json()
grid = None
if j.get("status") == "ok" and j.get("download_id"):
    grid = pd.read_excel(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
check("HTTP: fuzzy lookup ok + the typo row got its Dept in the download",
      j.get("status") == "ok" and grid is not None
      and grid.loc[grid["Name"] == "Jon Smith", "Dept"].iloc[0] == "Eng", str(j)[:150])

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
