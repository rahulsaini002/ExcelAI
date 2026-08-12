"""ENGINE PHASE 1.1 — Universal Formula Generator, Hands-layer verification (NO AI).

Runs hand-built formula plans through the REAL API (/inspect -> /execute -> /download)
and checks BOTH halves of every result:
  * the PREVIEW values pandas computed, and
  * the LIVE Excel formula openpyxl wrote into the downloaded workbook
plus M365 version warnings, spill behaviour, the div-zero guard, self-correction,
the phantom-column clarify, and honest redirects (GROUPBY/TEXTSPLIT/VLOOKUP).

Zero Brain calls: /execute runs caller-supplied plans; the two /process checks
monkeypatch llm.parse_instruction. Run from backend:
    .venv\\Scripts\\python.exe tests\\test_phase_1_1.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from datetime import date
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p11.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import llm  # noqa: E402
from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"
WB = TESTS / "standard_test_workbook.xlsx"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def small_wb() -> bytes:
    """Tiny deterministic sheet for date/financial checks."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Start", "End", "CashFlow", "Amount"])
    ws.append([date(2026, 1, 5), date(2026, 1, 12), -1000, 100])
    ws.append([date(2026, 2, 2), date(2026, 2, 16), 300, 250])
    ws.append([date(2026, 3, 2), date(2026, 3, 30), 420, 0])
    ws.append([date(2026, 4, 6), date(2026, 4, 20), 680, 50])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_formula(data: bytes, fname: str, name: str, formula: str):
    """inspect -> execute one add_formula_column -> (response json, downloaded workbook)."""
    sid = f"p11-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid}, files=[("files", (fname, data, OCT))])
    assert r.status_code == 200, r.text[:200]
    plan = json.dumps({"operations": [{"action": "add_formula_column", "name": name, "formula": formula}]})
    r = client.post("/execute", data={"session_id": sid, "plan": plan})
    j = r.json()
    book = None
    if j.get("status") == "ok" and j.get("download_id"):
        dl = client.get(f"/download/{j['download_id']}")
        book = openpyxl.load_workbook(io.BytesIO(dl.content), data_only=False)
    return j, book


def sheet_of(book, needle: str):
    name = next((s for s in book.sheetnames if needle in s), book.sheetnames[0])
    return book[name]


def col_idx(ws, header: str) -> int:
    return [c.value for c in ws[1]].index(header) + 1


def preview_col(j: dict, sheet_needle: str, column: str) -> list:
    """Column values from the response PREVIEW (the pandas-computed results). The saved
    file holds LIVE formulas without cached values, so read_excel would see blanks —
    the preview is where computed values are asserted."""
    for t in j.get("preview") or []:
        if sheet_needle in t.get("name", ""):
            return [row.get(column) for row in t.get("sample_rows", [])]
    t = (j.get("preview") or [{}])[0]
    return [row.get(column) for row in t.get("sample_rows", [])]


print("ENGINE PHASE 1.1 — Universal Formula Generator (Hands layer, no AI)\n")
wb_bytes = WB.read_bytes()
small = small_wb()

# ---- 1. IFS tiering: preview values + live formula --------------------------------------
j, book = run_formula(wb_bytes, "standard_test_workbook.xlsx",
                      "Tier", 'IFS({Qty}>20, "High", {Qty}>5, "Mid", TRUE, "Low")')
check("IFS executes (status ok)", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    ws = sheet_of(book, "Sales")
    f = ws.cell(row=2, column=col_idx(ws, "Tier")).value
    check("IFS live formula written with real refs", isinstance(f, str) and f.startswith("=IFS(")
          and "D2" in f, repr(f))
    check("IFS M365 note (Excel 2019)", "Excel 2019" in (j.get("explanation") or ""), j.get("explanation", "")[:160])

# ---- 2. XLOOKUP across sheets (messy-case keys!) -----------------------------------------
j, book = run_formula(wb_bytes, "standard_test_workbook.xlsx",
                      "Unit", 'XLOOKUP({Product}, {Prices.Product:}, {Prices.Unit_Price:}, 0)')
check("XLOOKUP executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    ws = sheet_of(book, "Sales")
    f = ws.cell(row=2, column=col_idx(ws, "Unit")).value
    check("XLOOKUP formula uses cross-sheet absolute ranges",
          isinstance(f, str) and "Prices!$A$2:$A$" in f and f.startswith("=XLOOKUP(C2"), repr(f))
    check("XLOOKUP M365 note (Excel 2021)", "Excel 2021" in (j.get("explanation") or ""))
    # messy-case products (widget / WIDGET / trailing spaces) still resolve in the preview
    vals = pd.to_numeric(pd.Series(preview_col(j, "Sales", "Unit")), errors="coerce").fillna(0)
    check("XLOOKUP preview resolves messy-case keys",
          len(vals) > 0 and (vals > 0).mean() > 0.9, f"sample={vals.head(4).tolist()}")

# ---- 3. SUMIF with a row criteria (per-region totals) ------------------------------------
j, book = run_formula(wb_bytes, "standard_test_workbook.xlsx",
                      "RegionTotal", "SUMIF({Region:}, {Region}, {Price:})")
check("SUMIF(row criteria) executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    ws = sheet_of(book, "Sales")
    f = (ws.cell(row=3, column=col_idx(ws, "RegionTotal")).value or "").replace(" ", "")
    check("SUMIF formula written with $ ranges + row ref",
          f.startswith("=SUMIF($B$2:$B$201,B3,$E$2:$E$201)"), repr(f))

# ---- 4. Text cleanup chain --------------------------------------------------------------
j, book = run_formula(wb_bytes, "standard_test_workbook.xlsx",
                      "CleanProduct", "PROPER(TRIM({Product}))")
check("PROPER(TRIM()) executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    got = [str(v) for v in preview_col(j, "Sales", "CleanProduct") if v is not None]
    check("preview values are trimmed + Proper Case",
          bool(got) and all(v == v.strip() and v[:1].isupper() and v[1:].islower() for v in got),
          f"sample={got[:4]}")

# ---- 5. Dates on a messy column: honest soft note ----------------------------------------
j, book = run_formula(wb_bytes, "standard_test_workbook.xlsx", "SaleYear", "YEAR({Date})")
check("YEAR({Date}) executes on messy dates", j.get("status") == "ok", j.get("error", "")[:140])
check("text-dates soft note present", "stored as text" in (j.get("explanation") or ""),
      j.get("explanation", "")[:180])

# ---- 6. NETWORKDAYS ----------------------------------------------------------------------
j, book = run_formula(small, "small.xlsx", "WorkDays", "NETWORKDAYS({Start}, {End})")
check("NETWORKDAYS executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    got = [int(v) for v in preview_col(j, "Data", "WorkDays") if v is not None]
    check("NETWORKDAYS values correct (incl. end date)", got == [6, 11, 21, 11], str(got))

# ---- 7. Financial ------------------------------------------------------------------------
j, _ = run_formula(small, "small.xlsx", "Payment", "PMT(0.01, 60, 500000)")
check("PMT executes", j.get("status") == "ok", j.get("error", "")[:140])
j, book = run_formula(small, "small.xlsx", "ProjectIRR", "IRR({CashFlow:})")
check("IRR executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    vals = [v for v in preview_col(j, "Data", "ProjectIRR") if v is not None]
    irr = float(vals[0]) if vals else float("nan")
    check("IRR value plausible (10-30% for -1000/300/420/680)", 0.10 < irr < 0.30, f"irr={irr}")

# ---- 8. Spill: UNIQUE written once ------------------------------------------------------
j, book = run_formula(wb_bytes, "standard_test_workbook.xlsx", "Regions", "UNIQUE({Region:})")
check("UNIQUE executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    ws = sheet_of(book, "Sales")
    c = col_idx(ws, "Regions")
    first, second = ws.cell(row=2, column=c).value, ws.cell(row=3, column=c).value
    check("spill formula written ONCE (row 2 only)",
          isinstance(first, str) and first.startswith("=UNIQUE($B$2:$B$201)") and second is None,
          f"row2={first!r} row3={second!r}")
    check("UNIQUE M365 note", "Microsoft 365" in (j.get("explanation") or ""))

# ---- 9. Guards ---------------------------------------------------------------------------
j, _ = run_formula(small, "small.xlsx", "Ratio", "{CashFlow} / {Amount}")
check("div-by-zero guarded (ok + blank note, not a crash)",
      j.get("status") == "ok" and "#DIV/0" in (j.get("explanation") or ""), j.get("explanation", "")[:160])

j, _ = run_formula(small, "small.xlsx", "Fixed", "{Amont} * 2")
check("typo with close match auto-repairs (Amont->Amount)",
      j.get("status") == "ok" and "Amount" in (j.get("explanation") or ""), j.get("explanation", "")[:160])

j, _ = run_formula(small, "small.xlsx", "Nope", "TEXTSPLIT({Amount}, \",\")")
check("TEXTSPLIT redirects honestly (no fake)",
      j.get("status") != "ok" and "split" in json.dumps(j).lower(), json.dumps(j)[:160])

j, _ = run_formula(small, "small.xlsx", "Nope2", "GROUPBY({Amount:}, {CashFlow:})")
check("GROUPBY redirects to pivot_summary",  # wording updated in Phase 2.1
      j.get("status") != "ok" and "pivot" in json.dumps(j).lower(), json.dumps(j)[:160])

j, _ = run_formula(small, "small.xlsx", "Nope3", "XLOOKUPP({Amount}, {Amount:}, {Amount:})")
check("unknown function gets did-you-mean",
      j.get("status") != "ok" and "XLOOKUP" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_formula(small, "small.xlsx", "Nope4", 'REGEXEXTRACT({Amount}, "([bad")')
check("invalid regex -> clean message", j.get("status") != "ok" and "regular expression" in json.dumps(j),
      json.dumps(j)[:160])

# ---- 10. Validator paths through /process (patched Brain, no API) ------------------------
real = llm.parse_instruction
try:
    llm.parse_instruction = lambda *a, **k: {
        "operations": [{"action": "add_formula_column", "name": "X",
                        "formula": "{Zorblatt} * 2"}], "confidence": 90}
    r = client.post("/process", data={"instruction": "x", "session_id": f"p11-{uuid.uuid4().hex[:8]}"},
                    files=[("files", ("small.xlsx", small, OCT))])
    j = r.json()
    check("phantom ref with NO close match -> clarify before execution",
          j.get("status") == "clarify" and "Zorblatt" in (j.get("clarification") or ""),
          json.dumps(j)[:160])

    llm.parse_instruction = lambda *a, **k: {
        "operations": [{"action": "add_formula_column", "name": "Y",
                        "formula": "XLOOKUP({Amount}, {Nope.Col:}, {Nope.Col:}, 0)"}], "confidence": 90}
    r = client.post("/process", data={"instruction": "x", "session_id": f"p11-{uuid.uuid4().hex[:8]}"},
                    files=[("files", ("small.xlsx", small, OCT))])
    j = r.json()
    check("missing sheet in {Sheet.Col:} -> clarify", j.get("status") == "clarify", json.dumps(j)[:160])
finally:
    llm.parse_instruction = real

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
