"""ENGINE PHASE 2.7 — print / page setup (NO AI).

Print setup rides on layout_format via a single free-text `print_setup` field (kept a
string, not a nested model, so the response schema stays under Gemini's serving limit).
This suite checks the plain-text parser, that every setting round-trips into the saved
.xlsx and reads back correctly (orientation, fit-to-page, print area, repeat title rows,
margins, header/footer), the title-row offset interaction, that the data is never
changed, and the honest failure paths.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_2_7.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p27.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import app.main as m  # noqa: E402
from app.db import init_db  # noqa: E402
from app.executor import execute_multi  # noqa: E402
from app.operations.base import OperationError  # noqa: E402
from app.operations.layout import layout_format, _parse_print_setup  # noqa: E402

init_db()
client = TestClient(m.app)
passed = failed = 0
OCT = "application/octet-stream"
DF = pd.DataFrame({"Region": ["N", "S", "E"], "Qty": [10, 20, 30], "Price": [100, 200, 150]})


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def saved(op: dict, df: pd.DataFrame = DF):
    """Run layout_format and return the saved worksheet (openpyxl) + the note."""
    _, note, d = layout_format(df, op)
    out, _, _ = m._serialize(df, "x.csv", "xlsx", [d])
    return load_workbook(io.BytesIO(out)).active, note


def err(op: dict, df: pd.DataFrame = DF) -> str:
    try:
        layout_format(df, op)
        return ""
    except OperationError as e:
        return str(e)


print("ENGINE PHASE 2.7 — print / page setup (no AI)\n")

# ---- (a) the plain-text parser ----
p = _parse_print_setup("make it landscape, fit to one page, repeat the header row, narrow margins")
check("parse: landscape + fit + repeat + narrow",
      p == {"orientation": "landscape", "fit_wide": 1, "repeat_header": True, "margins": "narrow"}, str(p))
check("parse: portrait", _parse_print_setup("portrait please").get("orientation") == "portrait")
check("parse: fit to N pages wide", _parse_print_setup("fit to 2 pages wide").get("fit_wide") == 2)
check("parse: print area range", _parse_print_setup("set print area to a1:f50").get("print_area") == "A1:F50")
check("parse: wide margins", _parse_print_setup("use wide margins").get("margins") == "wide")
check("parse: page numbers -> footer", _parse_print_setup("page numbers in the footer").get("footer_text") == "Page &P of &N")
check("parse: explicit footer text", _parse_print_setup("footer: Confidential").get("footer_text") == "Confidential")
check("parse: nothing recognizable -> empty", _parse_print_setup("do something nice") == {})

# ---- (b) each setting round-trips into the .xlsx ----
ws, note = saved({"print_setup": "landscape and fit to one page"})
check("xlsx: orientation landscape", ws.page_setup.orientation == "landscape", str(ws.page_setup.orientation))
check("xlsx: fit-to-page on (fitToWidth 1)",
      ws.page_setup.fitToWidth == 1 and ws.sheet_properties.pageSetUpPr
      and ws.sheet_properties.pageSetUpPr.fitToPage is True, "")
check("print note is plain language", "print setup" in note and "landscape" in note, note)

ws, _ = saved({"print_setup": "repeat the header row on every page"})
check("xlsx: repeat header rows (row 1, no title)", str(ws.print_title_rows) in ("1:1", "$1:$1"), str(ws.print_title_rows))

ws, _ = saved({"print_setup": "print area A1:B4"})  # openpyxl qualifies it as 'Sheet1'!$A$1:$B$4
check("xlsx: print area set", "A1:B4" in (ws.print_area or "").replace("$", ""), str(ws.print_area))

ws, _ = saved({"print_setup": "narrow margins"})
check("xlsx: narrow margins", abs(ws.page_margins.left - 0.25) < 1e-6, str(ws.page_margins.left))
ws, _ = saved({"print_setup": "wide margins"})
check("xlsx: wide margins", abs(ws.page_margins.left - 1.0) < 1e-6, str(ws.page_margins.left))

ws, _ = saved({"print_setup": "footer: Page &P of &N"})
check("xlsx: footer text", ws.oddFooter.center.text == "Page &P of &N", str(ws.oddFooter.center.text))
ws, _ = saved({"print_setup": "header: Q1 Report"})
check("xlsx: header text", ws.oddHeader.center.text == "Q1 Report", str(ws.oddHeader.center.text))

# ---- (c) interaction with a title row (offset): print area + repeat shift down ----
ws, _ = saved({"title": "Sales Report", "print_setup": "print area A1:C3, repeat the header row"})
check("title+print: print area shifts down with the inserted title row",
      "A2:C4" in (ws.print_area or "").replace("$", ""), str(ws.print_area))
check("title+print: repeat rows cover the title + header (1:2)",
      str(ws.print_title_rows) in ("1:2", "$1:$2"), str(ws.print_title_rows))
check("title+print: the title cell is actually there", ws["A1"].value == "Sales Report", str(ws["A1"].value))

# ---- (d) print setup can be the ONLY thing asked for (no other layout feature) ----
_, note, d = layout_format(DF, {"print_setup": "landscape"})
check("print-only request is valid (no other layout feature needed)", d["print"]["orientation"] == "landscape", str(d))

# ---- (e) data is never changed ----
res, _, notes, render = execute_multi({"t": DF.copy()}, "t", [{"action": "layout_format", "print_setup": "landscape, fit to one page"}])
check("print setup leaves the data unchanged", list(res.columns) == list(DF.columns) and len(res) == len(DF))
check("print setup emits a layout directive with print settings",
      any(r.get("type") == "layout" and r.get("print") for r in render), str(render))

# ---- (f) failures ----
check("unrecognized print request declined", "couldn't tell what print setting" in err({"print_setup": "do a barrel roll"}))
check("bad print area range declined", "isn't a print area" in err({"print_setup": "print area ZZ:top"})
      or "couldn't tell" in err({"print_setup": "print area ZZ:top"}))
check("empty layout (nothing at all) still declined", "what to change about the layout" in err({}))

# ---- (g) HTTP round-trip ----
buf = io.BytesIO()
with pd.ExcelWriter(buf, engine="openpyxl") as w:
    DF.to_excel(w, index=False, sheet_name="Sheet1")
sid = f"p27-{uuid.uuid4().hex[:10]}"
client.post("/inspect", data={"session_id": sid}, files=[("files", ("p.xlsx", buf.getvalue(), OCT))])
r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": [
    {"action": "layout_format", "print_setup": "landscape, fit to one page, repeat the header row"}]})})
j = r.json()
ws2 = None
if j.get("status") == "ok" and j.get("download_id"):
    ws2 = load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content)).active
check("HTTP: ok + print settings in the downloaded file",
      j.get("status") == "ok" and ws2 is not None and ws2.page_setup.orientation == "landscape"
      and ws2.page_setup.fitToWidth == 1 and str(ws2.print_title_rows) in ("1:1", "$1:$1"), str(j)[:150])

# ---- (h) regression guard: from_end still honored on a DIRECT plan (schema-swap) ----
res2, _, _, _ = execute_multi({"t": pd.DataFrame({"V": [1, 2, 3, 4, 5]})}, "t",
                              [{"action": "limit", "count": 2, "from_end": True}])
check("swap guard: 'keep last N' still works via a direct plan", list(res2["V"]) == [4, 5], str(list(res2["V"])))

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
