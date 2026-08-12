"""PROGRAM PHASE 0.2 — first operation, NO AI (PRD 1.4 sort / 1.13 download).

Proves the upload -> sort -> download round-trip through the real API with ZERO model
calls: /inspect seeds the session, /execute runs a hand-built Operation Plan (the
no-Brain path — exactly what the UI does after plan approval), /download streams the
result. No GEMINI_API_KEY needed.

Checked:
  1. round-trip completes (upload -> sort -> downloadable result)
  2. numbers sort NUMERICALLY — the 1,10,2,21,3 lexical trap, including numbers
     stored as text; row integrity kept (values stay paired with their row)
  3. descending variant
  4. the original is unchanged (upload bytes untouched; re-inspecting the original
     still shows the original order — nothing server-side mutated it)
  5. output opens cleanly in BOTH openpyxl and pandas (the "opens in Excel /
     Google Sheets" proxy), including from the messy 200-row Standard Workbook

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_0_2.py
(Needs tests/standard_test_workbook.xlsx — run make_standard_workbook.py first.)
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import sys
import tempfile
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p02.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def make_trap_xlsx() -> bytes:
    """Qty deliberately in the 1,10,2,21,3 order, with two values STORED AS TEXT —
    the classic lexical-sort trap (text sort gives 1,10,2,21,3; numeric 1,2,3,10,21)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Qty"])
    for name, qty in [("A", 1), ("B", "10"), ("C", 2), ("D", "21"), ("E", 3)]:
        ws.append([name, qty])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def inspect(data: bytes, name: str, sid: str):
    return client.post("/inspect", data={"session_id": sid},
                       files=[("files", (name, data, "application/octet-stream"))])


def execute_sort(sid: str, columns: list[str], order: str):
    plan = json.dumps({"operations": [{"action": "sort", "columns": columns, "orders": [order]}]})
    return client.post("/execute", data={"session_id": sid, "plan": plan})


def download(result_id: str) -> bytes:
    r = client.get(f"/download/{result_id}")
    assert r.status_code == 200, f"/download HTTP {r.status_code}"
    return r.content


print("PHASE 0.2 — sort round-trip, no AI (PRD 1.4 / 1.13)\n")

trap = make_trap_xlsx()
trap_sha = hashlib.sha256(trap).hexdigest()

# --- 1. upload + inspect shows the ORIGINAL (unsorted) order --------------------------
r = inspect(trap, "trap.xlsx", "p02-asc")
j = r.json()
sample = j["tables"][0]["sample_rows"] if r.status_code == 200 else []
orig_order = [row.get("Qty") for row in sample]
check("upload ok, preview shows original order (1,10,2,21,3)",
      r.status_code == 200 and [str(v) for v in orig_order] == ["1", "10", "2", "21", "3"],
      f"order={orig_order}")

# --- 2. execute the sort plan (NO model call) -----------------------------------------
r = execute_sort("p02-asc", ["Qty"], "asc")
j = r.json()
check("execute sort asc -> status ok, 5 rows",
      r.status_code == 200 and j.get("status") == "ok" and j.get("row_count") == 5,
      r.text[:160])
dl_id = j.get("download_id")
check("result offers a download id", bool(dl_id), f"keys={sorted(j.keys())[:12]}")

# --- 3. download opens cleanly + numeric order + row integrity ------------------------
out = download(dl_id)
wb = openpyxl.load_workbook(io.BytesIO(out))          # opens in openpyxl
df = pd.read_excel(io.BytesIO(out))                   # opens in pandas
qty = [int(v) for v in df["Qty"].tolist()]
names = df["Name"].tolist()
check("downloaded file opens cleanly (openpyxl + pandas)", bool(wb.sheetnames) and len(df) == 5)
check("NUMERIC sort: 1,2,3,10,21 — not the lexical 1,10,2,21,3",
      qty == [1, 2, 3, 10, 21], f"qty={qty}")
check("row integrity: names still paired with their Qty (A,C,E,B,D)",
      names == ["A", "C", "E", "B", "D"], f"names={names}")

# --- 4. descending variant -------------------------------------------------------------
r = execute_sort("p02-asc", ["Qty"], "desc")
j = r.json()
df_desc = pd.read_excel(io.BytesIO(download(j.get("download_id"))))
qty_desc = [int(v) for v in df_desc["Qty"].tolist()]
check("descending sort: 21,10,3,2,1", qty_desc == [21, 10, 3, 2, 1], f"qty={qty_desc}")

# --- 5. original unchanged --------------------------------------------------------------
check("original upload bytes untouched (hash identical)",
      hashlib.sha256(trap).hexdigest() == trap_sha)
r = inspect(trap, "trap.xlsx", "p02-fresh")
sample = r.json()["tables"][0]["sample_rows"] if r.status_code == 200 else []
check("re-inspecting the original still shows the ORIGINAL order",
      [str(row.get("Qty")) for row in sample] == ["1", "10", "2", "21", "3"],
      f"order={[row.get('Qty') for row in sample]}")

# --- 6. the messy Standard Workbook leg (200 rows, text prices, blanks) -----------------
wb_bytes = (TESTS / "standard_test_workbook.xlsx").read_bytes()
r = inspect(wb_bytes, "standard_test_workbook.xlsx", "p02-messy")
check("standard workbook inspects ok", r.status_code == 200, r.text[:120])
r = execute_sort("p02-messy", ["Price"], "asc")
j = r.json()
check("messy Sales sheet sorts by Price (status ok, 200 rows)",
      r.status_code == 200 and j.get("status") == "ok" and j.get("row_count") == 200,
      r.text[:160])
out = pd.read_excel(io.BytesIO(download(j.get("download_id"))), sheet_name=None)
sales_name = next((k for k in out if "Sales" in k), list(out.keys())[0])
prices = out[sales_name]["Price"].tolist()
numeric = [float(str(p).replace(",", "")) for p in prices
           if p is not None and str(p).strip() not in ("", "nan") and str(p).replace(",", "").replace(".", "", 1).lstrip("-").isdigit()]
check("messy Price column: numeric subsequence is non-decreasing (text-stored numbers included)",
      numeric == sorted(numeric), f"first 10: {numeric[:10]}")
check("messy output keeps all 200 rows", len(out[sales_name]) == 200, str(len(out[sales_name])))

# =====================================================================================
# The rest of the PRD 1.4 table (c–f) and 1.13-d, exercised directly against the
# executor — the round-trip above already proved the API path.
# =====================================================================================
from app.executor import execute_multi  # noqa: E402


def sorted_by(df, columns, orders):
    res, _, _, _ = execute_multi({"t": df.copy()}, "t",
                                 [{"action": "sort", "columns": columns, "orders": orders}])
    return res


# --- 1.4-c  dates sort CHRONOLOGICALLY, whatever the display format ------------------
# The Standard Workbook deliberately mixes real datetimes with date STRINGS.
dates = pd.DataFrame({
    "Date": [pd.Timestamp("2026-03-01"), "2026-01-15", pd.Timestamp("2026-02-10"), "2025-12-31"],
    "Tag": ["mar", "jan", "feb", "dec25"],
})
check("1.4-c dates sort chronologically across mixed formats",
      list(sorted_by(dates, ["Date"], ["asc"])["Tag"]) == ["dec25", "jan", "feb", "mar"],
      str(list(sorted_by(dates, ["Date"], ["asc"])["Tag"])))

# --- 1.4-d  text sorts alphabetically, ignoring case AND stray whitespace ------------
txt = pd.DataFrame({"R": ["banana", "Apple", "cherry", "APPLE", "Banana"]})
got = [s.lower() for s in sorted_by(txt, ["R"], ["asc"])["R"]]
check("1.4-d text sorts case-insensitively", got == sorted(got), str(got))

# Regression for the whitespace bug: ' East' / 'East  ' / 'EAST' must land together.
# Keying on raw text scattered them across the sheet, which defeats sorting entirely.
messy_sales = pd.read_excel(TESTS / "standard_test_workbook.xlsx", sheet_name="Sales")
norm = [str(v).strip().lower() for v in sorted_by(messy_sales, ["Region"], ["asc"])["Region"].dropna()]
runs = [k for i, k in enumerate(norm) if i == 0 or norm[i - 1] != k]
check("1.4-d each region forms ONE contiguous block despite case/space variants",
      len(runs) == len(set(runs)) and runs == sorted(runs), str(runs))
check("1.4-d sorting does NOT alter the stored values (padding preserved)",
      any(v != str(v).strip() for v in sorted_by(messy_sales, ["Region"], ["asc"])["Region"].dropna()),
      "values were trimmed — sorting must never rewrite data")

# --- 1.4-e  secondary sort ----------------------------------------------------------
sec = pd.DataFrame({"Region": ["N", "S", "N", "S"], "Qty": [5, 3, 9, 7]})
check("1.4-e secondary sort (Region asc, then Qty desc)",
      list(zip(sorted_by(sec, ["Region", "Qty"], ["asc", "desc"])["Region"],
               sorted_by(sec, ["Region", "Qty"], ["asc", "desc"])["Qty"]))
      == [("N", 9), ("N", 5), ("S", 7), ("S", 3)], "")

# --- 1.4-f  blanks grouped at the END, both directions -------------------------------
blanks = pd.DataFrame({"Qty": [5, None, 1, None, 3]})
for order in ("asc", "desc"):
    vals = list(sorted_by(blanks, ["Qty"], [order])["Qty"])
    check(f"1.4-f blanks sort to the end ({order})",
          all(pd.isna(v) for v in vals[-2:]) and not any(pd.isna(v) for v in vals[:3]), str(vals))

# --- 1.13-d  several operations in one session, each downloadable --------------------
sid = "p02-multi"
# Upload an .xlsx (not a .csv) so "opens cleanly in Excel" is a meaningful check — the
# engine deliberately returns a CSV for a CSV input, which openpyxl can't read.
_multi = io.BytesIO()
pd.DataFrame({"Name": ["A", "B", "A", "C"], "Qty": [3, 1, 3, 2]}).to_excel(_multi, index=False)
client.post("/inspect", data={"session_id": sid},
            files=[("files", ("m.xlsx", _multi.getvalue(),
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))])
step1 = client.post("/execute", data={"session_id": sid, "plan": json.dumps(
    {"operations": [{"action": "sort", "columns": ["Qty"], "orders": ["asc"]}]})}).json()
step2 = client.post("/execute", data={"session_id": sid, "plan": json.dumps(
    {"operations": [{"action": "remove_duplicates"}]})}).json()
check("1.13-d step 1 of a session returns a downloadable result",
      step1.get("status") == "ok" and step1.get("download_id"), str(step1)[:120])
check("1.13-d step 2 chains on step 1 and is separately downloadable",
      step2.get("status") == "ok" and step2.get("download_id")
      and step2["download_id"] != step1["download_id"] and step2.get("row_count") == 3,
      f"rows={step2.get('row_count')} (4 rows, one dup -> expect 3)")
for label, res in (("step 1", step1), ("step 2", step2)):
    blob = client.get(f"/download/{res['download_id']}").content
    try:
        openpyxl.load_workbook(io.BytesIO(blob))
        opens = True
    except Exception:
        opens = False
    check(f"1.13-d {label}'s file opens cleanly", opens, "")

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
