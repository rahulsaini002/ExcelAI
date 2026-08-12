"""ENGINE PHASE 2.3 — statistical analysis (NO AI).

describe / correlation / regression / moving_average / t_test, each checked for a
CORRECT result AND a plain-language interpretation, plus the honesty rails:
insufficient-data declines, text-column declines, and the pure-NumPy Student-t
p-value verified against known reference values (so regression/t-test significance
is trustworthy without SciPy).

Direct Hands calls check maths + messages; a few rows go through the real API
(/inspect → /execute → /download) to confirm result tables and the regression/t-test
summary sheets land in the workbook with the data left intact.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_2_3.py
"""
from __future__ import annotations

import io
import json
import math
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p23.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402
from app.operations.base import OperationError  # noqa: E402
from app.operations.statistics import statistics, student_t_two_sided_p  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def st(df: pd.DataFrame, op: dict):
    return statistics(df.copy(), op)


def err(df: pd.DataFrame, op: dict) -> str:
    try:
        statistics(df.copy(), op)
        return ""
    except OperationError as e:
        return str(e)


print("ENGINE PHASE 2.3 — statistical analysis (no AI)\n")

# ============ p-value accuracy (the one bit of custom maths) ============
# Reference two-sided p-values (standard t-tables / scipy.stats.t.sf*2).
for t, dfree, expect in [(2.0, 10, 0.0734), (2.228, 10, 0.0500), (1.0, 5, 0.3632),
                         (3.0, 20, 0.0071), (2.101, 18, 0.0500)]:
    p = student_t_two_sided_p(t, dfree)
    check(f"Student-t p(t={t}, df={dfree}) ≈ {expect}", abs(p - expect) < 0.003, f"got {p:.4f}")
check("Student-t p(0) == 1", student_t_two_sided_p(0.0, 8) == 1.0)

# ============ describe ============
PRICE = [10.0, 9.5, 10.5, 11.0, 10.0, 12.0, 8.0, 9.0]
UNITS = [100, 110, 90, 80, 100, 60, 130, 120]
DF = pd.DataFrame({"Price": PRICE, "Units": UNITS, "Region": list("NSNSNSNS"), "Label": list("ABCDEFGH")})

res, note, d = st(DF, {"stat_method": "describe"})
row = dict(zip(res["Statistic"], res["Price"]))
check("describe: count/mean/median correct",
      row["count"] == 8 and abs(row["mean"] - np.mean(PRICE)) < 1e-6
      and abs(row["median"] - np.median(PRICE)) < 1e-6, str(row))
check("describe: std is sample std (ddof=1)",
      abs(row["std"] - np.std(PRICE, ddof=1)) < 1e-6, f"{row['std']} vs {np.std(PRICE, ddof=1)}")
check("describe: range = max - min", abs(row["range"] - (max(PRICE) - min(PRICE))) < 1e-6, str(row["range"]))
check("describe: skips the text Label column", "Label" not in res.columns, str(list(res.columns)))
check("describe: plain-language note gives an example figure",
      "averages" in note and "Price" in note, note[:120])
check("describe: text-only data declined", "numeric" in err(pd.DataFrame({"X": list("abc")}), {"stat_method": "describe"}))

# numbers-stored-as-text are coerced, not skipped
res2, _, _ = st(pd.DataFrame({"V": ["10", "20", "30", "40"]}), {"stat_method": "describe"})
check("describe: numbers-stored-as-text coerced (mean 25)",
      abs(dict(zip(res2["Statistic"], res2["V"]))["mean"] - 25.0) < 1e-6, res2.to_string())

# ============ correlation ============
res, note, d = st(DF, {"stat_method": "correlation"})
# ground truth Pearson r(Price, Units)
truth = np.corrcoef(PRICE, UNITS)[0, 1]
rv = res[res.iloc[:, 0] == "Price"]["Units"].iloc[0]
check("correlation: matrix value matches numpy", abs(rv - truth) < 1e-3, f"{rv} vs {truth}")
check("correlation: diagonal is 1.0", res[res.iloc[:, 0] == "Price"]["Price"].iloc[0] == 1.0, res.to_string())
check("correlation: note names strongest pair + direction",
      "Price" in note and "Units" in note and ("negative" in note or "positive" in note), note[:160])
check("correlation: 'not causation' caveat present", "causation" in note.lower(), note[-80:])
check("correlation: needs 2+ varying numeric cols → declined",
      "two numeric" in err(pd.DataFrame({"A": [1, 2, 3], "B": list("xyz")}), {"stat_method": "correlation"}))
check("correlation: too few rows → declined",
      "at least 3" in err(pd.DataFrame({"A": [1, 2], "B": [3, 4]}), {"stat_method": "correlation"}))

# ============ regression ============
rng = np.random.default_rng(7)
x = np.arange(30, dtype=float)
y = 5.0 + 2.5 * x + rng.normal(0, 3, size=30)  # slope 2.5, intercept 5, strong fit
RDF = pd.DataFrame({"x": x, "y": y})
_, note, d = st(RDF, {"stat_method": "regression", "x_column": "x", "y_columns": ["y"]})
rows = {r[0]: r[1] for r in d["rows"]}
check("regression: recovers slope ≈ 2.5", abs(rows["Slope"] - 2.5) < 0.2, str(rows["Slope"]))
check("regression: recovers intercept ≈ 5", abs(rows["Intercept"] - 5.0) < 3.0, str(rows["Intercept"]))
check("regression: R² high (> 0.9) for a strong linear fit", rows["R-squared"] > 0.9, str(rows["R-squared"]))
check("regression: slope highly significant (p < 0.001)", float(rows["Slope p-value"]) < 0.001, str(rows["Slope p-value"]))
check("regression: note explains R² as % variance + significance",
      "%" in note and ("significant" in note), note[:200])
check("regression: writes a 'Regression' summary sheet + leaves data unchanged",
      d["type"] == "stats_sheet" and d["sheet_name"] == "Regression")

# a flat (no relationship) outcome → NOT significant, low R²
yflat = rng.normal(50, 10, size=30)
_, _, d2 = st(pd.DataFrame({"x": x, "y": yflat}), {"stat_method": "regression", "x_column": "x", "y_columns": ["y"]})
rows2 = {r[0]: r[1] for r in d2["rows"]}
check("regression: no real relationship → low R² and not significant",
      rows2["R-squared"] < 0.3 and float(rows2["Slope p-value"]) > 0.05,
      f"R2={rows2['R-squared']} p={rows2['Slope p-value']}")

check("regression: missing predictor/outcome declined",
      "predictor and an outcome" in err(RDF, {"stat_method": "regression"}))
check("regression: too few rows declined",
      "at least 3" in err(pd.DataFrame({"x": [1, 2], "y": [3, 4]}),
                          {"stat_method": "regression", "x_column": "x", "y_columns": ["y"]}))
check("regression: constant predictor declined",
      "same value" in err(pd.DataFrame({"x": [5, 5, 5, 5], "y": [1, 2, 3, 4]}),
                          {"stat_method": "regression", "x_column": "x", "y_columns": ["y"]}))

# ============ moving_average ============
mdf = pd.DataFrame({"Sales": [10, 20, 30, 40, 50, 60]})
res, note, d = st(mdf, {"stat_method": "moving_average", "column": "Sales", "count": 3})
ma = list(res["Sales_MA3"])
check("moving_average: adds MA column, first (window-1) blank",
      math.isnan(ma[0]) and math.isnan(ma[1]) and ma[2] == 20.0 and ma[3] == 30.0, str(ma))
check("moving_average: note explains the leading blanks + purpose",
      "blank" in note and "smooth" in note.lower(), note[:120])
check("moving_average: window > rows declined",
      "larger than the table" in err(mdf, {"stat_method": "moving_average", "column": "Sales", "count": 99}))
check("moving_average: window < 2 declined",
      "at least 2" in err(mdf, {"stat_method": "moving_average", "column": "Sales", "count": 1}))
check("moving_average: text column declined",
      "looks like text" in err(pd.DataFrame({"C": list("abcd")}),
                               {"stat_method": "moving_average", "column": "C", "count": 2}))
# safe fallback: Brain omits the column, but there's exactly ONE numeric column -> use it
res, _, _ = st(pd.DataFrame({"Sales": [10, 20, 30, 40], "City": list("abcd")}),
               {"stat_method": "moving_average", "count": 2})
check("moving_average: no column + single numeric col -> uses it",
      "Sales_MA2" in res.columns, str(list(res.columns)))
check("moving_average: no column + several numeric cols -> asks (never guesses)",
      "Which column" in err(pd.DataFrame({"A": [1, 2, 3], "B": [4, 5, 6]}),
                            {"stat_method": "moving_average", "count": 2}))

# ============ t_test ============
# two numeric columns, clearly different means
a = [10, 12, 11, 13, 10, 12, 11, 13]
b = [20, 22, 21, 23, 20, 22, 21, 23]
_, note, d = st(pd.DataFrame({"A": a, "B": b}), {"stat_method": "t_test", "columns": ["A", "B"]})
rows = {r[0]: r[1] for r in d["rows"]}
check("t_test: big mean gap → significant (p < 0.05)", float(rows["p-value"]) < 0.05, str(rows["p-value"]))
check("t_test: note states both means + verdict",
      "mean" in note and ("significant" in note), note[:160])
check("t_test: writes a 'T-Test' summary sheet", d["sheet_name"] == "T-Test")

# same distribution → NOT significant
same = pd.DataFrame({"A": [10, 11, 12, 13, 10, 11], "B": [11, 10, 13, 12, 11, 10]})
_, _, d2 = st(same, {"stat_method": "t_test", "columns": ["A", "B"]})
check("t_test: equal-ish groups → not significant",
      float({r[0]: r[1] for r in d2["rows"]}["p-value"]) > 0.05, str(d2["rows"]))

# value column split by a 2-group column
gdf = pd.DataFrame({"Sales": a + b, "Region": ["North"] * 8 + ["South"] * 8})
_, note, d = st(gdf, {"stat_method": "t_test", "value_column": "Sales", "group_by": ["Region"]})
check("t_test: value+group mode compares the two groups",
      "North" in note and "South" in note, note[:160])
check("t_test: not exactly 2 groups declined",
      "exactly TWO" in err(pd.DataFrame({"S": [1, 2, 3], "G": ["a", "b", "c"]}),
                           {"stat_method": "t_test", "value_column": "S", "group_by": ["G"]}))
check("t_test: a group with <2 values declined",
      "at least 2" in err(pd.DataFrame({"S": [1, 2, 3], "G": ["a", "b", "b"]}),
                          {"stat_method": "t_test", "value_column": "S", "group_by": ["G"]}))
check("t_test: no columns/groups given declined",
      "two numeric columns" in err(pd.DataFrame({"A": [1, 2], "B": [3, 4], "C": [5, 6], "D": [7, 8]}),
                                   {"stat_method": "t_test"}))
# safe fallback: Brain omits columns, but there are EXACTLY two numeric columns -> compare them
_, note2, _ = st(pd.DataFrame({"A": a, "B": b, "Name": list("abcdefgh")}), {"stat_method": "t_test"})
check("t_test: no columns + exactly two numeric cols -> compares them",
      "'A'" in note2 and "'B'" in note2, note2[:120])

# ============ dispatch + empty ============
check("unknown stat_method declined with the menu",
      "I can describe" in err(DF, {"stat_method": "anova"}))
check("empty table declined", "no data" in err(pd.DataFrame(columns=["A"]), {"stat_method": "describe"}))

# ============ HTTP round-trip ============
def api(df: pd.DataFrame, op: dict):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Sheet1")
    sid = f"p23-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", ("stats.xlsx", buf.getvalue(), OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": [op]})})
    j = r.json()
    book = None
    if j.get("status") == "ok" and j.get("download_id"):
        book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
    return j, book

j, book = api(RDF, {"action": "statistics", "stat_method": "regression", "x_column": "x", "y_columns": ["y"]})
check("HTTP regression: ok + 'Regression' sheet + original data sheet intact",
      j.get("status") == "ok" and book is not None and "Regression" in book.sheetnames
      and book["Sheet1"]["A1"].value == "x", str(j)[:160] if book is None else str(book.sheetnames))

j, book = api(DF, {"action": "statistics", "stat_method": "describe"})
check("HTTP describe: downloads a stats table",
      book is not None and book.active["A1"].value == "Statistic", str(j)[:160])

j, book = api(mdf, {"action": "statistics", "stat_method": "moving_average", "column": "Sales", "count": 3})
check("HTTP moving_average: adds the MA column to the download",
      book is not None and any((book.active.cell(row=1, column=c).value or "").startswith("Sales_MA")
                               for c in range(1, 4)), str(j)[:160])

j, _ = api(pd.DataFrame({"A": [1, 2]}), {"action": "statistics", "stat_method": "correlation"})
check("HTTP correlation on too-little data: friendly decline, not a 500",
      j.get("status") != "ok" and "at least" in json.dumps(j).lower(), str(j)[:160])

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
