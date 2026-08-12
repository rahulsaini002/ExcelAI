"""ENGINE PHASE 1.7 — native Excel Tables, Hands-layer verification (NO AI).

Hand-built excel_table plans through the real API, then openpyxl readback of ws.tables:
style + banding, refs, the live =SUBTOTAL() totals row (auto and explicit), name
sanitizing, the replace-on-rerun guard, the empty-sheet decline, and the Phase-1.4
title-shift interplay.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_1_7.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p17.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def sales_xlsx(rows: int = 4) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Product", "Qty", "Price"])
    data = [["North", "Widget", 5, 100.0], ["South", "Gadget", 25, 200.0],
            ["East", "Pin", 10, 50.0], ["West", "Gizmo", 2, 400.0]]
    for r in data[:rows]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_ops(ops: list[dict], data: bytes | None = None):
    sid = f"p17-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", ("sales.xlsx", data or sales_xlsx(), OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": ops})})
    j = r.json()
    ws = None
    if j.get("status") == "ok" and j.get("download_id"):
        book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
        ws = book[book.sheetnames[0]]
    return j, ws


print("ENGINE PHASE 1.7 — native Excel Tables (Hands layer, no AI)\n")

# ---- basic table ---------------------------------------------------------------------------
j, ws = run_ops([{"action": "excel_table"}])
check("plain 'format as table' executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    tables = list(ws.tables.values())
    check("one Table object, correct ref (A1:D5)", len(tables) == 1 and str(tables[0].ref) == "A1:D5",
          f"{[(t.displayName, str(t.ref)) for t in tables]}")
    check("default blue style + banded rows", tables[0].tableStyleInfo.name == "TableStyleMedium2"
          and tables[0].tableStyleInfo.showRowStripes, str(tables[0].tableStyleInfo))

# ---- styled ---------------------------------------------------------------------------------
j, ws = run_ops([{"action": "excel_table", "table_style": "green"}])
check("green style maps to Medium7", j.get("status") == "ok" and ws is not None
      and list(ws.tables.values())[0].tableStyleInfo.name == "TableStyleMedium7",
      j.get("error", "")[:140])

# ---- totals: auto ---------------------------------------------------------------------------
j, ws = run_ops([{"action": "excel_table", "totals": True}])
check("auto totals row executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    t = list(ws.tables.values())[0]
    check("ref extended for totals (A1:D6) + totalsRowCount", str(t.ref) == "A1:D6"
          and t.totalsRowCount == 1, f"ref={t.ref} trc={t.totalsRowCount}")
    check("live SUBTOTAL(109) on numeric columns", ws["C6"].value == "=SUBTOTAL(109,C2:C5)"
          and ws["D6"].value == "=SUBTOTAL(109,D2:D5)", f"C6={ws['C6'].value!r} D6={ws['D6'].value!r}")
    check("'Total' label in the first text column", ws["A6"].value == "Total", repr(ws["A6"].value))
    check("note names the totals", "totals row" in (j.get("explanation") or ""),
          j.get("explanation", "")[:160])

# ---- totals: explicit spec ------------------------------------------------------------------
j, ws = run_ops([{"action": "excel_table", "totals_spec": [{"column": "Qty", "agg": "average"}]}])
check("explicit average totals executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    check("SUBTOTAL(101) average on Qty only", ws["C6"].value == "=SUBTOTAL(101,C2:C5)"
          and ws["D6"].value is None, f"C6={ws['C6'].value!r} D6={ws['D6'].value!r}")

# ---- name sanitizing -----------------------------------------------------------------------
j, ws = run_ops([{"action": "excel_table", "table_name": "My Sales Table!"}])
check("name sanitized (no spaces/punct)", j.get("status") == "ok" and ws is not None
      and list(ws.tables.values())[0].displayName == "My_Sales_Table_",
      f"{[t for t in (ws.tables if ws else {})]}")

# ---- replace on rerun (two table ops in one plan) --------------------------------------------
j, ws = run_ops([{"action": "excel_table", "table_style": "blue"},
                 {"action": "excel_table", "table_style": "green"}])
check("re-running replaces, never corrupts (ONE table, latest style)",
      j.get("status") == "ok" and ws is not None and len(ws.tables) == 1
      and list(ws.tables.values())[0].tableStyleInfo.name == "TableStyleMedium7",
      f"{[(t.displayName, t.tableStyleInfo.name) for t in (ws.tables.values() if ws else [])]}")

# ---- title-shift interplay -------------------------------------------------------------------
j, ws = run_ops([{"action": "excel_table", "totals": True},
                 {"action": "layout_format", "title": "Sales Book"}])
check("table + title plan executes", j.get("status") == "ok", j.get("error", "")[:160])
if ws:
    t = list(ws.tables.values())[0]
    check("table ref SHIFTED below the title (A2:D7)", str(t.ref) == "A2:D7", str(t.ref))
    check("totals formula shifted too", ws["C7"].value == "=SUBTOTAL(109,C3:C6)", repr(ws["C7"].value))
    check("title in A1", ws["A1"].value == "Sales Book", repr(ws["A1"].value))

# ---- failure paths ---------------------------------------------------------------------------
empty = openpyxl.Workbook()
empty.active.title = "Sales"
empty.active.append(["Region", "Qty"])
buf = io.BytesIO()
empty.save(buf)
j, _ = run_ops([{"action": "excel_table"}], data=buf.getvalue())
check("empty sheet (headers only) -> declined honestly", j.get("status") != "ok"
      and "at least one" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_ops([{"action": "excel_table", "table_style": "polkadot"}])
check("unknown style -> clean error naming styles", j.get("status") != "ok"
      and "blue" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_ops([{"action": "excel_table", "totals_spec": [{"column": "Nope", "agg": "sum"}]}])
check("totals on a missing column -> clean error", j.get("status") != "ok"
      and "Nope" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_ops([{"action": "excel_table", "totals_spec": [{"column": "Qty", "agg": "median"}]}])
check("unsupported agg -> clean error listing aggs", j.get("status") != "ok"
      and "average" in json.dumps(j), json.dumps(j)[:160])

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
