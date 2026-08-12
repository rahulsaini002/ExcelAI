"""Generate the extra Phase-0.1 fixture files (Build & Test Program, Phase 0.1).

The Standard Test Workbook covers multi-sheet/mess; these cover the PRD 1.1 rows it
can't: a plain CSV, blank/duplicate headers, a mixed-type column, an empty file, a
wrong file type, a corrupt xlsx, a large file, non-English headers, and a real PDF
(borrowed from pdfplumber's bundled test assets — PDF ingestion is a SUPPORTED
input since engine Phase 4.1, superseding the program doc's older "wrong type
(.pdf)" line; the wrong-type case uses .txt instead).

Run from backend:  .venv\\Scripts\\python.exe tests\\make_extra_files.py
Writes into:       tests/files/
"""
from __future__ import annotations

import csv
import io
import pathlib
import random
from datetime import date

from openpyxl import Workbook

FILES = pathlib.Path(__file__).resolve().parent / "files"
FILES.mkdir(exist_ok=True)
rng = random.Random(7)


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build() -> None:
    # normal, clean single-sheet xlsx (the "happy path" control)
    (FILES / "normal.xlsx").write_bytes(xlsx({
        "Sheet1": [["Date", "Item", "Qty", "Price"],
                   [date(2026, 1, 5), "Pen", 2, 10.5],
                   [date(2026, 1, 6), "Book", 1, 99.0],
                   [date(2026, 1, 7), "Bag", 3, 450.0]],
    }))

    # plain CSV
    with open(FILES / "basic.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Region", "Product", "Qty", "Price"])
        for i in range(25):
            w.writerow([rng.choice(["North", "South"]), rng.choice(["Pen", "Book"]),
                        rng.randrange(1, 20), round(rng.uniform(10, 500), 2)])

    # blank header in the middle
    (FILES / "blank_header.xlsx").write_bytes(xlsx({
        "Sheet1": [["Name", None, "Amount"], ["A", "x", 1], ["B", "y", 2]],
    }))

    # duplicate headers
    (FILES / "dup_headers.xlsx").write_bytes(xlsx({
        "Sheet1": [["Amount", "Amount", "Name"], [1, 2, "A"], [3, 4, "B"]],
    }))

    # one column mixing numbers, text, and dates
    (FILES / "mixed_types.xlsx").write_bytes(xlsx({
        "Sheet1": [["Id", "Value"], [1, 42], [2, "not-a-number"],
                   [3, date(2026, 2, 1)], [4, 7.5], [5, "N/A"]],
    }))

    # empty file (zero bytes) + wrong type + corrupt xlsx
    (FILES / "empty.xlsx").write_bytes(b"")
    (FILES / "wrong_type.txt").write_text("just some prose, definitely not a table\n", encoding="utf-8")
    (FILES / "corrupt.xlsx").write_bytes(b"PK\x03\x04 this used to be a workbook" + bytes(rng.randrange(256) for _ in range(512)))

    # large: 60,001 data rows (crosses the engine's 50k "Large file" note threshold)
    with open(FILES / "large.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["N", "Group", "Value"])
        for i in range(60_001):
            w.writerow([i, "G" + str(i % 7), i * 3 % 1000])

    # non-English (Hindi + Urdu) headers, UTF-8
    (FILES / "non_english.xlsx").write_bytes(xlsx({
        "Sheet1": [["नाम", "राशि", "تاریخ"], ["अशोक", 1200, "2026-01-05"],
                   ["Sana", 850, "2026-01-06"], ["राहुल", 2100, "2026-01-07"]],
    }))

    # Phase 0.3 (Brain) fixtures:
    # column-name approximation target — "sort by rev" must resolve to "Revenue"
    (FILES / "approx.xlsx").write_bytes(xlsx({
        "Sheet1": [["Region", "Revenue", "Units"],
                   ["North", 4200, 12], ["South", 8800, 31], ["East", 1500, 7]],
    }))
    # genuinely ambiguous column — "sort by price" matches two columns equally
    (FILES / "ambiguous.xlsx").write_bytes(xlsx({
        "Sheet1": [["Product", "Price_2024", "Price_2025"],
                   ["Pen", 10, 12], ["Book", 99, 105], ["Bag", 450, 430]],
    }))

    # Phase 1.3 (split/merge/fill-by-example) fixture: two-part names to split, city/state
    # to merge, and a Code column with NO delimiter (the ambiguous-split failure case).
    (FILES / "names.xlsx").write_bytes(xlsx({
        "Sheet1": [["Full Name", "City", "State", "Code"],
                   ["Asha Sharma", "Pune", "MH", "X9K2A1"],
                   ["Rahul Verma", "Jaipur", "RJ", "B7Q4Z8"],
                   ["Meera", "Kochi", "KL", "M3N5P0"],
                   ["Sana Ali Khan", "Delhi", "DL", "T6R1W9"]],
    }))

    # Phase 2.2 (reshaping) fixtures: WIDE monthly data to unpivot, LONG data to pivot,
    # and a small metrics table to transpose.
    (FILES / "reshape_wide.xlsx").write_bytes(xlsx({
        "Sheet1": [["Region", "Jan", "Feb", "Mar"],
                   ["North", 10, 30, 50], ["South", 20, 40, 60], ["East", 5, 15, 25]],
    }))
    (FILES / "reshape_long.xlsx").write_bytes(xlsx({
        "Sheet1": [["Region", "Month", "Sales"],
                   ["North", "Jan", 10], ["North", "Jan", 5], ["South", "Jan", 20],
                   ["North", "Feb", 7], ["South", "Mar", 12]],  # ragged: not every combo present
    }))
    (FILES / "reshape_metrics.xlsx").write_bytes(xlsx({
        "Sheet1": [["Metric", "Q1", "Q2", "Q3"],
                   ["Revenue", 100, 200, 150], ["Cost", 40, 60, 55], ["Profit", 60, 140, 95]],
    }))

    # Phase 2.3 (statistics) fixture: a TINY table (2 rows) for the insufficient-data
    # honesty row — a regression/correlation here must decline, never invent a fit.
    (FILES / "stats_tiny.xlsx").write_bytes(xlsx({
        "Sheet1": [["Qty", "Price"], [3, 100], [5, 220]],
    }))

    # Phase 3.1 (fuzzy lookup) fixture: an Orders sheet with a name typo ('Jon Smith')
    # that must fuzzy-match the Roster's 'John Smith'.
    (FILES / "fuzzy_pair.xlsx").write_bytes(xlsx({
        "Orders": [["Name", "Order"], ["John Smith", 1], ["Jon Smith", 2],
                   ["Mary Jones", 3], ["Asha Rao", 4]],
        "Roster": [["FullName", "Dept"], ["John Smith", "Eng"], ["Mary Jones", "Sales"],
                   ["Asha Rao", "Ops"]],
    }))

    # Phase 2.9 (workbook compare) fixture: two sheets that differ (Rahul's Amount
    # changed, ID 4→5 swapped, a Status column added) — for "what changed?" rows.
    (FILES / "compare_pair.xlsx").write_bytes(xlsx({
        "Before": [["ID", "Name", "Amount"], [1, "Asha", 100], [2, "Rahul", 200],
                   [3, "Meera", 150], [4, "Vikram", 300]],
        "After": [["ID", "Name", "Amount", "Status"], [1, "Asha", 100, "ok"],
                  [2, "Rahul", 250, "ok"], [3, "Meera", 150, "ok"], [5, "Dev", 400, "new"]],
    }))

    # Phase 1.10: legacy input formats — .ods via pandas' odf engine; .xls via xlwt
    # (xlwt is a dev-only fixture helper: pandas can READ .xls with xlrd but no longer
    # writes it).
    import pandas as _pd

    ldf = _pd.DataFrame({"Region": ["North", "South", "East"], "Amount": [120, 340, 95]})
    ldf.to_excel(FILES / "legacy.ods", index=False, engine="odf")
    try:
        import xlwt

        wb_xls = xlwt.Workbook()
        sh = wb_xls.add_sheet("Old")
        for c, h in enumerate(ldf.columns):
            sh.write(0, c, h)
        for r, row in enumerate(ldf.itertuples(index=False), start=1):
            for c, v in enumerate(row):
                sh.write(r, c, v)
        wb_xls.save(str(FILES / "legacy.xls"))
    except ImportError:
        print("NOTE: xlwt not installed — legacy.xls fixture skipped")

    # a real PDF: prefer pdfplumber's bundled samples; else draw one with PyMuPDF —
    # a RULED 4x3 grid (lines + cell text) so pdfplumber's line-based table
    # extraction can find it digitally, without needing OCR.
    pdf_note = "skipped (no generator available)"
    try:
        import pdfplumber
        pkg = pathlib.Path(pdfplumber.__file__).parent
        candidates = list(pkg.rglob("*table*.pdf")) + list(pkg.rglob("*.pdf"))
        if candidates:
            (FILES / "sample.pdf").write_bytes(candidates[0].read_bytes())
            pdf_note = f"copied from {candidates[0].name}"
        else:
            raise FileNotFoundError
    except Exception:
        try:
            import fitz  # PyMuPDF

            doc = fitz.open()
            page = doc.new_page()
            rows = [["Name", "Qty", "Price"], ["Pen", "2", "10.50"],
                    ["Book", "1", "99.00"], ["Bag", "3", "450.00"]]
            x0, y0, cw, rh = 72, 72, 120, 24
            for ri in range(len(rows) + 1):          # horizontal rules
                y = y0 + ri * rh
                page.draw_line((x0, y), (x0 + cw * 3, y))
            for ci in range(4):                      # vertical rules
                x = x0 + ci * cw
                page.draw_line((x, y0), (x, y0 + rh * len(rows)))
            for ri, row in enumerate(rows):          # cell text
                for ci, val in enumerate(row):
                    page.insert_text((x0 + ci * cw + 6, y0 + ri * rh + 16), val, fontsize=10)
            doc.save(FILES / "sample.pdf")
            doc.close()
            pdf_note = "generated with PyMuPDF (ruled 4x3 table)"
        except Exception as exc:
            pdf_note = f"skipped ({exc})"

    # An image that is NOT a readable table (PRD 1.1-h). Sumio accepts photos of tables via
    # OCR (Phase 4.1), so the honest behaviour for an unreadable one is a clear 4xx telling
    # the user what to upload — never an opaque 500. Regression file for that fix.
    (FILES / "photo.jpg").write_bytes(b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01" + b"not a real image" * 8)

    for p in sorted(FILES.iterdir()):
        print(f"  {p.name:20} {p.stat().st_size:>9,} bytes")
    print(f"PDF: {pdf_note}")


if __name__ == "__main__":
    build()
