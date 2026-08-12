"""ENGINE PHASE 1.4 — layout & formatting polish, Hands-layer verification (NO AI).

Freeze panes, computed autofit, borders, header fill, Indian lakh/crore number format,
and the title row — including the hard case: a plan that writes a LIVE FORMULA and a
CF RULE and THEN adds a title row must shift both down correctly (openpyxl Translator
for formulas, sqref rewrite for CF).

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_1_4.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p14.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def sales_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Product", "Qty", "Price", "AmountLongColumn"])
    ws.append(["North", "Widget", 5, 100.5, 1234567.0])
    ws.append(["South", "A very long product name here", 25, 200.0, 12345678.0])
    ws.append(["East", "Pin", 10, 50.0, 123456789.0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_plan(ops: list[dict]):
    sid = f"p14-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", ("sales.xlsx", sales_xlsx(), OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": ops})})
    j = r.json()
    ws = None
    if j.get("status") == "ok" and j.get("download_id"):
        book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
        ws = book[next((s for s in book.sheetnames if "Sales" in s), book.sheetnames[0])]
    return j, ws


L = lambda ws, header, row: ws.cell(row=row, column=[c.value for c in ws[1]].index(header) + 1)  # noqa: E731

print("ENGINE PHASE 1.4 — layout & formatting polish (Hands layer, no AI)\n")

# ---- freeze variants ----------------------------------------------------------------------
for spec, want in [("header", "A2"), ("first_column", "B1"), ("both", "B2"), ("C5", "C5")]:
    j, ws = run_plan([{"action": "layout_format", "freeze": spec}])
    check(f"freeze '{spec}' -> panes at {want}",
          j.get("status") == "ok" and ws is not None and ws.freeze_panes == want,
          f"got {ws.freeze_panes if ws else j.get('error', '')[:100]}")

# ---- autofit ------------------------------------------------------------------------------
j, ws = run_plan([{"action": "layout_format", "autofit": True}])
check("autofit executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    w_qty = ws.column_dimensions["C"].width
    w_prod = ws.column_dimensions["B"].width
    check("autofit: long-content column wider than short one",
          w_prod and w_qty and w_prod > w_qty and w_prod <= 60, f"prod={w_prod} qty={w_qty}")

# ---- borders ------------------------------------------------------------------------------
j, ws = run_plan([{"action": "layout_format", "borders": "all"}])
check("borders 'all' executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    mid = ws.cell(row=3, column=3)
    check("grid border on an interior cell", mid.border.left.style == "thin"
          and mid.border.bottom.style == "thin", str(mid.border))

j, ws = run_plan([{"action": "layout_format", "borders": "outline"}])
if ws:
    corner = ws.cell(row=1, column=1)
    inner = ws.cell(row=2, column=3)
    check("outline: box edges only (corner yes, interior no)",
          corner.border.top.style == "thin" and corner.border.left.style == "thin"
          and inner.border.left.style is None, f"corner={corner.border} inner={inner.border}")

# ---- header fill --------------------------------------------------------------------------
j, ws = run_plan([{"action": "layout_format", "header_fill": "blue"}])
check("header fill executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    h = ws.cell(row=1, column=1)
    check("header tinted + bold", (h.fill.start_color.rgb or "").endswith("BDD7EE") and h.font.bold,
          f"fill={h.fill.start_color.rgb} bold={h.font.bold}")

# ---- Indian lakh/crore number format (via format_cells) -----------------------------------
j, ws = run_plan([{"action": "format_cells", "format_columns": ["AmountLongColumn"],
                   "number_format": "indian_currency", "decimals": 0}])
check("indian_currency format executes", j.get("status") == "ok", j.get("error", "")[:140])
if ws:
    code = L(ws, "AmountLongColumn", 2).number_format
    check("lakh/crore grouping code applied (##\\,## sections + conditions)",
          "[>=10000000]" in code and "#\\,##\\,##" in code, code[:80])

# ---- THE BIG ONE: formula + CF + title => everything shifts --------------------------------
j, ws = run_plan([
    {"action": "add_formula_column", "name": "Total", "formula": "{Qty} * {Price}"},
    {"action": "conditional_format", "columns": ["Price"], "rule_type": "greater_than",
     "value": 90, "color": "green"},
    {"action": "layout_format", "title": "Q1 Sales Report", "freeze": "header"},
])
check("combined formula+CF+title plan executes", j.get("status") == "ok", j.get("error", "")[:200])
if ws:
    check("title in merged A1 row", ws["A1"].value == "Q1 Sales Report"
          and any(str(r) == "A1:F1" for r in ws.merged_cells.ranges),
          f"A1={ws['A1'].value!r} merges={[str(r) for r in ws.merged_cells.ranges]}")
    heads = [c.value for c in ws[2]]
    check("headers moved to row 2", heads[:2] == ["Region", "Product"], str(heads))
    f = (ws.cell(row=3, column=heads.index("Total") + 1).value or "").replace(" ", "")
    check("live formula SHIFTED with the title (=C3*D3)", f == "=C3*D3", repr(f))
    cf = [(str(c.sqref), r.type) for c in ws.conditional_formatting for r in c.rules]
    check("CF range shifted down (D3:D5)", any(s == "D3:D5" and t == "cellIs" for s, t in cf), str(cf))
    check("freeze respects the title offset (A3)", ws.freeze_panes == "A3", str(ws.freeze_panes))

# ---- merge_range guard ---------------------------------------------------------------------
j, ws = run_plan([{"action": "layout_format", "merge_range": "A10:D10"}])
check("merge of a blank range executes", j.get("status") == "ok"
      and ws is not None and any(str(r) == "A10:D10" for r in ws.merged_cells.ranges),
      j.get("error", "")[:140])

j, _ = run_plan([{"action": "layout_format", "merge_range": "A1:D1"}])
check("merge over the HEADERS -> refused, points to title",
      j.get("status") != "ok" and "title" in json.dumps(j), json.dumps(j)[:180])

j, _ = run_plan([{"action": "layout_format", "merge_range": "A2:C3"}])
check("merge over DATA -> refused with the column named",
      j.get("status") != "ok" and "wipe data" in json.dumps(j), json.dumps(j)[:180])

# ---- failure paths --------------------------------------------------------------------------
j, _ = run_plan([{"action": "layout_format", "freeze": "the moon"}])
check("bad freeze value -> clean error", j.get("status") != "ok" and "freeze" in json.dumps(j),
      json.dumps(j)[:160])

j, _ = run_plan([{"action": "layout_format", "borders": "dotted-rainbow"}])
check("bad borders value -> clean error", j.get("status") != "ok" and "outline" in json.dumps(j),
      json.dumps(j)[:160])

j, _ = run_plan([{"action": "layout_format", "header_fill": "chartreuse"}])
check("bad color -> clean error naming palette", j.get("status") != "ok" and "green" in json.dumps(j),
      json.dumps(j)[:160])

j, _ = run_plan([{"action": "layout_format"}])
check("empty op -> asks what to change", j.get("status") != "ok" and "freeze" in json.dumps(j),
      json.dumps(j)[:160])

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
