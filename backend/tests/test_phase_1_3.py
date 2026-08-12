"""ENGINE PHASE 1.3 — split / merge / fill-by-example, Hands-layer verification (NO AI).

Hand-built plans through the real API (/inspect -> /execute -> /download): split by
delimiter / fixed width / regex (with honest missing-part counts), merge with a live
TEXTJOIN formula in the saved file, and deterministic fill-by-example induction
(usernames, initials, prefix codes) that REFUSES conflicting examples.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_1_3.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p13.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def people_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["Full Name", "City", "State", "Code"])
    ws.append(["Asha Sharma", "Pune", "MH", "AB123"])
    ws.append(["Rahul Verma", "Jaipur", "RJ", "CD456"])
    ws.append(["Meera", "Kochi", "KL", "EF789"])          # one-part name (split edge)
    ws.append(["Sana Ali Khan", "Delhi", "", "GH012"])    # three parts + blank State
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_ops(ops: list[dict]):
    sid = f"p13-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", ("people.xlsx", people_xlsx(), OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": ops})})
    return r.json()


def sample(j: dict, column: str) -> list:
    t = (j.get("preview") or [{}])[0]
    return [row.get(column) for row in t.get("sample_rows", [])]


print("ENGINE PHASE 1.3 — split / merge / fill-by-example (Hands layer, no AI)\n")

# ---- (a) split ---------------------------------------------------------------------------
j = run_ops([{"action": "split_column", "column": "Full Name",
              "new_columns": ["First", "Last"], "delimiter": " "}])
check("split by space executes", j.get("status") == "ok", j.get("error", "")[:140])
check("split values correct (First)", sample(j, "First")[:4] == ["Asha", "Rahul", "Meera", "Sana"],
      str(sample(j, "First")))
check("split values correct (Last; extra parts stay in last)",
      sample(j, "Last")[:4] == ["Sharma", "Verma", None, "Ali Khan"], str(sample(j, "Last")))
check("one-part row reported honestly", "1 row didn't have all the parts" in (j.get("explanation") or ""),
      j.get("explanation", "")[:160])

j = run_ops([{"action": "split_column", "column": "Full Name",
              "new_columns": ["First", "Last"]}])  # no delimiter -> inferred
check("delimiter inferred (space) when obvious", j.get("status") == "ok"
      and sample(j, "First")[0] == "Asha", j.get("error", "")[:140])

j = run_ops([{"action": "split_column", "column": "Code",
              "new_columns": ["Letters", "Digits"], "pattern": r"([A-Z]+)(\d+)"}])
check("regex-pattern split (capture groups)", j.get("status") == "ok"
      and sample(j, "Letters")[:2] == ["AB", "CD"] and sample(j, "Digits")[:2] == ["123", "456"],
      f"{sample(j, 'Letters')} {sample(j, 'Digits')}")

j = run_ops([{"action": "split_column", "column": "Code",
              "new_columns": ["Alpha", "Rest"], "widths": [2, 3]}])
check("fixed-width split", j.get("status") == "ok" and sample(j, "Alpha")[:2] == ["AB", "CD"],
      str(sample(j, "Alpha")))

j = run_ops([{"action": "split_column", "column": "Code", "new_columns": ["A", "B"], "delimiter": "|"}])
check("delimiter absent everywhere -> asks for the delimiter... or reports honestly",
      j.get("status") != "ok" or "didn't have all the parts" in (j.get("explanation") or ""),
      json.dumps(j)[:160])

j = run_ops([{"action": "split_column", "column": "Full Name", "new_columns": ["City", "Last"], "delimiter": " "}])
check("name collision -> clean error", j.get("status") != "ok" and "already exist" in json.dumps(j),
      json.dumps(j)[:160])

# ---- (b) merge ---------------------------------------------------------------------------
j = run_ops([{"action": "merge_columns", "columns": ["City", "State"],
              "name": "City_State", "separator": ", "}])
check("merge executes", j.get("status") == "ok", j.get("error", "")[:140])
check("merge values joined", sample(j, "City_State")[:2] == ["Pune, MH", "Jaipur, RJ"],
      str(sample(j, "City_State")))
check("blank part skipped (no dangling comma)", sample(j, "City_State")[3] == "Delhi",
      str(sample(j, "City_State")))
if j.get("download_id"):
    book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
    ws = book[book.sheetnames[0]]
    heads = [c.value for c in ws[1]]
    f = ws.cell(row=2, column=heads.index("City_State") + 1).value
    check("saved file has a LIVE TEXTJOIN formula",
          isinstance(f, str) and f.startswith('=TEXTJOIN(", ", TRUE,'), repr(f))

j = run_ops([{"action": "merge_columns", "columns": ["City", "State"], "name": "Where",
              "separator": " - ", "keep_original": False}])
check("merge with keep_original=false drops sources", j.get("status") == "ok"
      and sample(j, "City") == [None] * len(sample(j, "City")), str(sample(j, "City"))[:80])

j = run_ops([{"action": "merge_columns", "columns": ["City"], "name": "X"}])
check("merge with one column -> clean error", j.get("status") != "ok"
      and "at least two" in json.dumps(j), json.dumps(j)[:160])

# ---- (c) fill-by-example ------------------------------------------------------------------
j = run_ops([{"action": "fill_by_example", "column": "Full Name", "name": "Username",
              "examples": [{"input": "Asha Sharma", "output": "asha.sharma"},
                           {"input": "Rahul Verma", "output": "rahul.verma"}]}])
check("username pattern induced from two examples", j.get("status") == "ok"
      and sample(j, "Username")[:2] == ["asha.sharma", "rahul.verma"], str(sample(j, "Username")))
check("pattern applied to unseen rows", sample(j, "Username")[3] == "sana.ali",
      str(sample(j, "Username")))

j = run_ops([{"action": "fill_by_example", "column": "Full Name", "name": "Initials",
              "examples": [{"input": "Asha Sharma", "output": "A.S."}]}])
check("initials pattern (prefixes) from ONE example", j.get("status") == "ok"
      and sample(j, "Initials")[:2] == ["A.S.", "R.V."], str(sample(j, "Initials")))
check("single-example note recommends a check",
      "single example" in (j.get("explanation") or ""), j.get("explanation", "")[:160])

j = run_ops([{"action": "fill_by_example", "column": "City", "name": "CityCode",
              "examples": [{"input": "Pune", "output": "PUN"}]}])
check("prefix+upper code pattern", j.get("status") == "ok"
      and sample(j, "CityCode")[:3] == ["PUN", "JAI", "KOC"], str(sample(j, "CityCode")))

j = run_ops([{"action": "fill_by_example", "column": "Full Name", "name": "Bad",
              "examples": [{"input": "Asha Sharma", "output": "asha.sharma"},
                           {"input": "Rahul Verma", "output": "verma-RAHUL"}]}])
check("conflicting examples -> refuses with a clear ask", j.get("status") != "ok"
      and "same rule" in json.dumps(j), json.dumps(j)[:200])

j = run_ops([{"action": "fill_by_example", "column": "Full Name", "name": "Bad2",
              "examples": [{"input": "Asha Sharma", "output": "42banana!"}]}])
check("output unrelated to input -> refuses honestly", j.get("status") != "ok"
      and "re-check" in json.dumps(j), json.dumps(j)[:200])

j = run_ops([{"action": "fill_by_example", "column": "Full Name", "name": "Nope", "examples": []}])
check("no examples -> clean error", j.get("status") != "ok" and "at least one example" in json.dumps(j),
      json.dumps(j)[:160])

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
