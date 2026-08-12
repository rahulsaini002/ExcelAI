"""ENGINE PHASE 1.6 — sheet management, Hands-layer verification (NO AI).

Hand-built sheet_op plans through the real API, with openpyxl readback of the saved
workbook: new-tab results, rename (order preserved), delete (with the only-sheet
guard), copy, move, tab colors, hide/unhide — plus the put-the-summary-in-a-new-tab
end-to-end flow and every failure path.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_1_6.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p16.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"
WB = (TESTS / "standard_test_workbook.xlsx").read_bytes()


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def one_sheet() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Only"
    ws.append(["A"])
    ws.append([1])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_ops(ops: list[dict], data: bytes = WB, fname: str = "standard_test_workbook.xlsx"):
    sid = f"p16-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", (fname, data, OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": ops})})
    j = r.json()
    book = None
    if j.get("status") == "ok" and j.get("download_id"):
        book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
    return j, book


def tails(book) -> list[str]:
    return [s.split(" - ")[-1] for s in book.sheetnames]


print("ENGINE PHASE 1.6 — sheet management (Hands layer, no AI)\n")

# ---- put the summary in a new tab (the flagship flow) --------------------------------------
j, book = run_ops([
    {"action": "aggregate", "group_by": ["Region"], "agg_column": "Price", "agg_func": "sum"},
    {"action": "sheet_op", "sheet_action": "new_sheet", "new_name": "Report"},
])
check("aggregate -> new tab 'Report' executes", j.get("status") == "ok", j.get("error", "")[:160])
if book:
    check("workbook keeps ALL tabs + Report", "Report" in tails(book) and "Customers" in tails(book),
          str(book.sheetnames))
    rep = book[next(s for s in book.sheetnames if s.endswith("Report"))]
    check("Report holds the summary (Region column, grouped rows — 26 messy variants)",
          rep.cell(row=1, column=1).value == "Region" and rep.max_row < 30,
          f"head={rep.cell(row=1, column=1).value} rows={rep.max_row}")

# ---- rename (incl. order preserved) ---------------------------------------------------------
j, book = run_ops([{"action": "sheet_op", "sheet_action": "rename",
                    "sheet_name": "Sales", "new_name": "Raw"}])
check("rename Sales -> Raw executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    t = tails(book)
    check("renamed, order preserved (Raw first)", t[0] == "Raw" and "Sales" not in t, str(t))

# ---- delete + guards -----------------------------------------------------------------------
j, book = run_ops([{"action": "sheet_op", "sheet_action": "delete", "sheet_name": "Empty"}])
check("delete Empty executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    check("Empty gone, others kept", "Empty" not in tails(book) and "Sales" in tails(book),
          str(book.sheetnames))

j, _ = run_ops([{"action": "sheet_op", "sheet_action": "delete"}],  # working = the only one
               data=one_sheet(), fname="single.xlsx")
check("delete the ONLY sheet -> declined", j.get("status") != "ok"
      and "only sheet" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_ops([{"action": "sheet_op", "sheet_action": "delete", "sheet_name": "Nonexistent"}])
check("delete a missing sheet -> clean error listing sheets",
      j.get("status") != "ok" and "Available" in json.dumps(j), json.dumps(j)[:160])

# ---- copy ----------------------------------------------------------------------------------
j, book = run_ops([{"action": "sheet_op", "sheet_action": "copy", "sheet_name": "Prices"}])
check("copy Prices executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    t = tails(book)
    check("copy lands right after its source", "Prices Copy" in t
          and t.index("Prices Copy") == t.index("Prices") + 1, str(t))
    src = book[next(s for s in book.sheetnames if s.endswith("Prices"))]
    cpy = book[next(s for s in book.sheetnames if s.endswith("Prices Copy"))]
    check("copy has the same data", src.max_row == cpy.max_row
          and src.cell(row=2, column=1).value == cpy.cell(row=2, column=1).value,
          f"{src.max_row} vs {cpy.max_row}")

# ---- move ----------------------------------------------------------------------------------
j, book = run_ops([{"action": "sheet_op", "sheet_action": "move",
                    "sheet_name": "Prices", "position": "first"}])
check("move Prices first executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    check("Prices is the first tab", tails(book)[0] == "Prices", str(tails(book)))

j, _ = run_ops([{"action": "sheet_op", "sheet_action": "move", "sheet_name": "Prices"}])
check("move without a position -> clean ask", j.get("status") != "ok"
      and "first" in json.dumps(j), json.dumps(j)[:160])

# ---- tab color / hide / unhide --------------------------------------------------------------
j, book = run_ops([{"action": "sheet_op", "sheet_action": "tab_color",
                    "sheet_name": "Prices", "tab_color": "green"}])
check("tab color executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    ws = book[next(s for s in book.sheetnames if s.endswith("Prices"))]
    got = ws.sheet_properties.tabColor
    check("Prices tab is green in the saved file", got is not None and "63BE7B" in str(got.rgb or got.value),
          str(got))

j, book = run_ops([{"action": "sheet_op", "sheet_action": "hide", "sheet_name": "Empty"}])
check("hide executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    ws = book[next(s for s in book.sheetnames if s.endswith("Empty"))]
    check("Empty is hidden in the saved file", ws.sheet_state == "hidden", ws.sheet_state)

j, _ = run_ops([{"action": "sheet_op", "sheet_action": "hide"}],  # working = the only one
               data=one_sheet(), fname="single.xlsx")
check("hide the only sheet -> declined", j.get("status") != "ok"
      and "visible" in json.dumps(j), json.dumps(j)[:160])

# ---- more failure paths ---------------------------------------------------------------------
j, _ = run_ops([{"action": "sheet_op", "sheet_action": "rename", "sheet_name": "Sales",
                 "new_name": "Prices"}])
check("rename onto an existing name -> clean error", j.get("status") != "ok"
      and "already exists" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_ops([{"action": "sheet_op", "sheet_action": "new_sheet"}])
check("new sheet without a name -> clean ask", j.get("status") != "ok"
      and "called" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_ops([{"action": "sheet_op", "sheet_action": "teleport", "sheet_name": "Sales"}])
check("unknown sheet action -> clean error", j.get("status") != "ok"
      and "rename" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_ops([{"action": "sheet_op", "sheet_action": "tab_color", "sheet_name": "Prices",
                 "tab_color": "polkadot"}])
check("unknown tab color -> clean error naming colors", j.get("status") != "ok"
      and "green" in json.dumps(j), json.dumps(j)[:160])

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
