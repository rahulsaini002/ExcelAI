"""ENGINE PHASE 1.5 — data validation / dropdowns, Hands-layer verification (NO AI).

Hand-built data_validation plans through the real API, then openpyxl readback of
ws.data_validations: dropdowns (explicit, derived-from-column, long-list helper sheet),
number/date/text-length ranges, custom {Col} rules, prompts/alerts, the violation-count
honesty note, the title-shift interplay, and the failure paths.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_1_5.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from datetime import date
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p15.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def orders_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Region", "Qty", "Date", "Notes"])
    ws.append(["North", 5, date(2026, 2, 1), "ok"])
    ws.append(["South", 2500, date(2025, 12, 30), "x" * 30])   # Qty + Date violations
    ws.append(["north ", 10, date(2026, 6, 15), ""])
    ws.append(["Fantasyland", 50, date(2026, 3, 3), "hello"])  # list violation
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_dv(ops: list[dict]):
    sid = f"p15-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", ("orders.xlsx", orders_xlsx(), OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": ops})})
    j = r.json()
    book = None
    if j.get("status") == "ok" and j.get("download_id"):
        book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
    return j, book


def dvs(book):
    name = next((s for s in book.sheetnames if "Orders" in s),
                next(s for s in book.sheetnames if "Options" not in s))
    ws = book[name]
    return ws, list(ws.data_validations.dataValidation)


print("ENGINE PHASE 1.5 — data validation / dropdowns (Hands layer, no AI)\n")

# ---- explicit dropdown --------------------------------------------------------------------
j, book = run_dv([{"action": "data_validation", "columns": ["Region"], "validation_type": "list",
                   "allowed_values": ["North", "South", "East", "West"],
                   "input_message": "Pick a region", "error_message": "Regions only"}])
check("explicit dropdown executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    ws, rules = dvs(book)
    check("list rule on A2:A5 with inline options", len(rules) == 1 and rules[0].type == "list"
          and rules[0].formula1 == '"North,South,East,West"' and "A2:A5" in str(rules[0].sqref),
          f"{[(r.type, r.formula1, str(r.sqref)) for r in rules]}")
    check("prompt + error alert set", rules[0].prompt == "Pick a region" and rules[0].error == "Regions only",
          f"{rules[0].prompt!r} {rules[0].error!r}")
    check("violation count in note (Fantasyland + case/space variants are counted honestly)",
          "existing value" in (j.get("explanation") or ""), j.get("explanation", "")[:180])

# ---- derived dropdown ----------------------------------------------------------------------
j, book = run_dv([{"action": "data_validation", "columns": ["Region"], "validation_type": "list"}])
check("derived dropdown executes (from the column's own values)", j.get("status") == "ok"
      and "from the column's own values" in (j.get("explanation") or ""), j.get("error", "")[:140])
if book:
    _, rules = dvs(book)
    check("derived options include the real values", rules and "North" in rules[0].formula1
          and "Fantasyland" in rules[0].formula1, rules[0].formula1 if rules else "none")

# ---- long list -> hidden helper sheet ------------------------------------------------------
many = [f"Option {i:03d}" for i in range(40)]  # 40 x ~10 chars > 250 inline cap
j, book = run_dv([{"action": "data_validation", "columns": ["Region"], "validation_type": "list",
                   "allowed_values": many}])
check("long list executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    _, rules = dvs(book)
    helper = [s for s in book.sheetnames if "Options" in s]
    check("long list -> hidden helper sheet + range ref", bool(helper)
          and book[helper[0]].sheet_state == "hidden" and rules
          and rules[0].formula1.startswith("=") and "$A$1:$A$40" in rules[0].formula1,
          f"sheets={book.sheetnames} f1={rules[0].formula1 if rules else None}")

# ---- whole-number range --------------------------------------------------------------------
j, book = run_dv([{"action": "data_validation", "columns": ["Qty"], "validation_type": "whole",
                   "min_value": 1, "max_value": 1000}])
check("Qty 1-1000 executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    _, rules = dvs(book)
    check("whole/between rule with bounds", rules and rules[0].type == "whole"
          and rules[0].operator == "between" and rules[0].formula1 == "1" and rules[0].formula2 == "1000",
          f"{[(r.type, r.operator, r.formula1, r.formula2) for r in rules]}")
    check("existing violation (2500) counted", "1 existing value" in (j.get("explanation") or ""),
          j.get("explanation", "")[:180])

# ---- date range ----------------------------------------------------------------------------
j, book = run_dv([{"action": "data_validation", "columns": ["Date"], "validation_type": "date",
                   "min_value": "2026-01-01", "max_value": "2026-12-31"}])
check("2026-only date rule executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    _, rules = dvs(book)
    check("date rule uses DATE() bounds", rules and rules[0].type == "date"
          and rules[0].formula1 == "DATE(2026,1,1)" and rules[0].formula2 == "DATE(2026,12,31)",
          f"{[(r.type, r.formula1, r.formula2) for r in rules]}")
    check("2025 date counted as existing violation", "1 existing value" in (j.get("explanation") or ""),
          j.get("explanation", "")[:180])

# ---- text length ---------------------------------------------------------------------------
j, book = run_dv([{"action": "data_validation", "columns": ["Notes"], "validation_type": "text_length",
                   "max_value": 20}])
check("text-length rule executes + violation counted", j.get("status") == "ok"
      and "1 existing value" in (j.get("explanation") or ""), j.get("explanation", "")[:180])

# ---- custom formula ------------------------------------------------------------------------
j, book = run_dv([{"action": "data_validation", "columns": ["Qty"], "validation_type": "custom",
                   "formula": "{Qty} < 100000"}])
check("custom {Col} rule executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    _, rules = dvs(book)
    check("custom formula rendered with real refs", rules and rules[0].type == "custom"
          and rules[0].formula1 == "$B2 < 100000", f"{rules[0].formula1 if rules else None}")

# ---- interplay with the Phase-1.4 title shift ----------------------------------------------
j, book = run_dv([
    {"action": "data_validation", "columns": ["Qty"], "validation_type": "whole",
     "min_value": 1, "max_value": 1000},
    {"action": "layout_format", "title": "Order Book"},
])
check("dv + title plan executes", j.get("status") == "ok", j.get("error", "")[:140])
if book:
    ws, rules = dvs(book)
    check("dv range SHIFTED below the title (B3:B6)", rules and "B3:B6" in str(rules[0].sqref),
          f"{[str(r.sqref) for r in rules]}")
    check("title landed in A1", ws["A1"].value == "Order Book", repr(ws["A1"].value))

# ---- failure paths --------------------------------------------------------------------------
j, _ = run_dv([{"action": "data_validation", "columns": ["Salary"], "validation_type": "whole",
                "min_value": 1}])
check("non-existent column -> clean error naming it", j.get("status") != "ok" and "Salary" in json.dumps(j),
      json.dumps(j)[:160])

j, _ = run_dv([{"action": "data_validation", "columns": ["Qty"], "validation_type": "whole",
                "min_value": 1000, "max_value": 1}])
check("min > max -> clean error", j.get("status") != "ok" and "swap" in json.dumps(j), json.dumps(j)[:160])

j, _ = run_dv([{"action": "data_validation", "columns": ["Date"], "validation_type": "date",
                "min_value": "someday"}])
check("unreadable date -> clean error", j.get("status") != "ok" and "start date" in json.dumps(j),
      json.dumps(j)[:160])

j, _ = run_dv([{"action": "data_validation", "columns": ["Qty"], "validation_type": "psychic"}])
check("unknown type -> clean error listing options", j.get("status") != "ok" and "dropdown" in json.dumps(j),
      json.dumps(j)[:160])

j, _ = run_dv([{"action": "data_validation", "columns": ["Notes"], "validation_type": "whole"}])
check("range rule without bounds -> clean error", j.get("status") != "ok" and "minimum" in json.dumps(j),
      json.dumps(j)[:160])

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
