"""Generate the Standard Test Workbook — the deliberately messy .xlsx every test
phase reuses (Build & Test Program, "Global test assets", asset A).

One workbook, five sheets:
  Sales     ~200 rows of Date/Region/Product/Qty/Price with REAL mess: exact duplicate
            rows, blank cells, numbers stored as text, four different date formats
            (some as true dates, some as strings), mixed-case region/product names,
            and trailing/leading spaces.
  Customers Customer_ID/Name/Email with case+space variants and duplicate emails.
  Prices    Product/Unit_Price — the lookup key sheet (product names match Sales
            only modulo case/space, on purpose).
  Empty     completely empty (no header, no rows).
  SingleRow header + exactly one data row.

Deterministic: seeded RNG, so the file is byte-for-byte reproducible content-wise —
tests can rely on exact row counts and known dirt.

Run from backend:  .venv\\Scripts\\python.exe tests\\make_standard_workbook.py
Writes:            tests/standard_test_workbook.xlsx
"""
from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parent / "standard_test_workbook.xlsx"
rng = random.Random(42)

REGIONS_CLEAN = ["North", "South", "East", "West"]
PRODUCTS_CLEAN = ["Widget", "Gadget", "Doohickey", "Gizmo", "Sprocket"]


def messy_text(clean: str) -> str:
    """Return the value in one of several dirty spellings (or clean)."""
    styles = [
        lambda s: s,                      # clean
        lambda s: s.lower(),              # lowercase
        lambda s: s.upper(),              # UPPERCASE
        lambda s: s + "  ",               # trailing spaces
        lambda s: " " + s,                # leading space
        lambda s: s.lower() + " ",        # lowercase + trailing
    ]
    return rng.choice(styles)(clean)


def messy_date(d: date):
    """Return the date as a true date cell OR one of three string formats."""
    styles = [
        lambda x: x,                                        # real date cell
        lambda x: x.isoformat(),                            # "2026-01-15"
        lambda x: x.strftime("%d/%m/%Y"),                   # "15/01/2026"
        lambda x: x.strftime("%b %d, %Y"),                  # "Jan 15, 2026"
    ]
    return rng.choice(styles)(d)


def build() -> None:
    wb = Workbook()

    # ------------------------------------------------------------------ Sales ----
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Date", "Region", "Product", "Qty", "Price"])

    base = date(2026, 1, 1)
    rows: list[list] = []
    for _ in range(180):
        d = base + timedelta(days=rng.randrange(0, 150))
        region = messy_text(rng.choice(REGIONS_CLEAN))
        product = messy_text(rng.choice(PRODUCTS_CLEAN))
        qty: object = rng.randrange(1, 50)
        price: object = round(rng.uniform(99.0, 4999.0), 2)

        # numbers stored as text (~1 in 6)
        if rng.random() < 0.17:
            qty = str(qty)
        if rng.random() < 0.17:
            price = f"{price}"

        # blank cells (~1 in 12 per blankable column)
        if rng.random() < 0.08:
            region = None
        if rng.random() < 0.08:
            qty = None
        if rng.random() < 0.08:
            price = None

        rows.append([messy_date(d), region, product, qty, price])

    # exact duplicate rows: re-append 20 copies of random earlier rows
    for _ in range(20):
        rows.append(list(rng.choice(rows[:180])))

    for r in rows:
        ws.append(r)

    # -------------------------------------------------------------- Customers ----
    ws = wb.create_sheet("Customers")
    ws.append(["Customer_ID", "Name", "Email"])
    first = ["Asha", "Rahul", "Meera", "Vikram", "Sana", "John", "Priya", "Dev"]
    last = ["Sharma", "Patel", "Khan", "Iyer", "Das", "Kaur"]
    emails: list[str] = []
    for i in range(1, 31):
        name = f"{rng.choice(first)} {rng.choice(last)}"
        email = f"{name.split()[0].lower()}.{name.split()[1].lower()}{i}@example.com"
        emails.append(email)
        # case/space variants on name + email
        ws.append([f"C{i:03d}", messy_text(name), messy_text(email) if rng.random() < 0.4 else email])
    # a few duplicate emails (same address, different id/case)
    for j, dup in enumerate(rng.sample(emails, 4)):
        ws.append([f"C9{j:02d}", "Duplicate Person", dup.upper() if j % 2 else dup + " "])

    # ----------------------------------------------------------------- Prices ----
    ws = wb.create_sheet("Prices")
    ws.append(["Product", "Unit_Price"])
    for p in PRODUCTS_CLEAN:
        ws.append([p, round(rng.uniform(80.0, 4500.0), 2)])

    # ------------------------------------------------------------------ Empty ----
    wb.create_sheet("Empty")  # deliberately: no header, no rows

    # -------------------------------------------------------------- SingleRow ----
    ws = wb.create_sheet("SingleRow")
    ws.append(["Item", "Value"])
    ws.append(["only-row", 1])

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"  Sales rows: {len(rows)} (incl. 20 exact duplicates)")
    print("  Sheets: Sales, Customers, Prices, Empty, SingleRow")


if __name__ == "__main__":
    build()
