"""ENGINE PHASE 1.9 — explain changes as cell notes, Hands-layer verification (NO AI).

Hand-built plans ending in explain_changes, with openpyxl readback of the Comments in
the saved workbook: fill -> notes on exactly the filled cells; dedupe -> A1 summary
(and NO per-cell notes — row positions shifted); added column -> header note; the
title-shift interplay; data immutability; and the no-op/first-op gentle paths.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_1_9.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p19.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import init_db  # noqa: E402
from app.main import app  # noqa: E402

init_db()
client = TestClient(app)
passed = failed = 0
OCT = "application/octet-stream"


def check(name: str, cond: bool, detail: str = "") -> None:
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def gaps_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Qty"])
    ws.append(["North", 5])
    ws.append(["South", None])   # blank -> will be filled (sheet row 3)
    ws.append(["North", 5])      # duplicate of row 2
    ws.append(["East", None])    # blank -> will be filled (sheet row 5)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_ops(ops: list[dict]):
    sid = f"p19-{uuid.uuid4().hex[:10]}"
    r = client.post("/inspect", data={"session_id": sid},
                    files=[("files", ("gaps.xlsx", gaps_xlsx(), OCT))])
    assert r.status_code == 200, r.text[:200]
    r = client.post("/execute", data={"session_id": sid, "plan": json.dumps({"operations": ops})})
    j = r.json()
    ws = None
    if j.get("status") == "ok" and j.get("download_id"):
        book = openpyxl.load_workbook(io.BytesIO(client.get(f"/download/{j['download_id']}").content))
        ws = book[book.sheetnames[0]]
    return j, ws


def commented(ws) -> dict[str, str]:
    out = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.comment is not None:
                out[cell.coordinate] = cell.comment.text
    return out


print("ENGINE PHASE 1.9 — explain changes as cell notes (Hands layer, no AI)\n")

# ---- fill -> notes on exactly the filled cells ----------------------------------------------
j, ws = run_ops([
    {"action": "fill_missing", "columns": ["Qty"], "fill_value": "0"},
    {"action": "explain_changes"},
])
check("fill + explain executes", j.get("status") == "ok", j.get("error", "")[:160])
if ws:
    notes = commented(ws)
    check("notes on exactly the two filled cells (B3, B5)",
          set(notes) == {"B3", "B5"}, str(notes))
    check("note text says was-blank -> now 0", "blank" in notes.get("B3", "")
          and "0" in notes.get("B3", ""), notes.get("B3", ""))
    check("notes authored by Sumio", notes.get("B3", "").startswith("Sumio:"), notes.get("B3", ""))
    check("underlying data unchanged (values intact)", ws["B3"].value in (0, "0")
          and ws["A2"].value == "North", f"B3={ws['B3'].value!r}")
    check("response note explains where to look", "Hover" in (j.get("explanation") or ""),
          j.get("explanation", "")[:180])

# ---- dedupe -> A1 summary, NO per-cell notes -------------------------------------------------
j, ws = run_ops([
    {"action": "remove_duplicates"},
    {"action": "explain_changes"},
])
check("dedupe + explain executes", j.get("status") == "ok", j.get("error", "")[:160])
if ws:
    notes = commented(ws)
    check("summary note on A1 only (rows shifted -> per-cell would lie)",
          set(notes) == {"A1"}, str(set(notes)))
    check("A1 note counts the removed row", "1 row removed" in notes.get("A1", ""),
          notes.get("A1", "")[:160])
    check("honest response note about skipped per-cell", "skipped" in (j.get("explanation") or ""),
          j.get("explanation", "")[:200])

# ---- added column -> header note -------------------------------------------------------------
j, ws = run_ops([
    {"action": "add_formula_column", "name": "Double", "formula": "{Qty} * 2"},
    {"action": "explain_changes"},
])
check("formula column + explain executes", j.get("status") == "ok", j.get("error", "")[:160])
if ws:
    notes = commented(ws)
    heads = [c.value for c in ws[1]]
    head_cell = f"{openpyxl.utils.get_column_letter(heads.index('Double') + 1)}1"
    check("header note on the added column", head_cell in notes
          and "Double" in notes[head_cell], f"{head_cell} in {set(notes)}")

# ---- title-shift interplay -------------------------------------------------------------------
j, ws = run_ops([
    {"action": "fill_missing", "columns": ["Qty"], "fill_value": "0"},
    {"action": "explain_changes"},
    {"action": "layout_format", "title": "Filled Data"},
])
check("explain + title plan executes", j.get("status") == "ok", j.get("error", "")[:160])
if ws:
    notes = commented(ws)
    check("notes SHIFTED below the title (B4, B6)", set(notes) == {"B4", "B6"}, str(set(notes)))
    check("title present", ws["A1"].value == "Filled Data", repr(ws["A1"].value))

# ---- gentle paths ----------------------------------------------------------------------------
j, _ = run_ops([{"action": "explain_changes"}])
check("explain as the ONLY op -> gentle guidance (no crash)", j.get("status") in ("ok", "message")
      and "add it after" in json.dumps(j), json.dumps(j)[:200])

j, _ = run_ops([
    {"action": "fill_missing", "columns": ["Region"], "fill_value": "X"},  # no blanks in Region
    {"action": "explain_changes"},
])
check("nothing actually changed -> honest 'nothing to annotate'",
      j.get("status") == "ok" and "nothing" in (j.get("explanation") or "").lower(),
      j.get("explanation", "")[:200])

print(f"\n{passed} passed, {failed} failed.")
try:
    os.unlink(_db)
except Exception:
    pass
sys.exit(1 if failed else 0)
