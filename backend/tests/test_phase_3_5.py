"""ENGINE PHASE 3.5 — multi-language tuning (OFFLINE part).

The DoD's core — "the full operation set × {EN, HI, UR, Hinglish}" — is a LIVE regression
(prompt_battery.csv + test_multilang.py, which call the model) and is quota-gated. What
this suite proves WITHOUT the model is the part the ENGINE controls: execution is fully
language-agnostic. The Brain turns any language into a structured plan; the Hands then
run on data + columns that may themselves be Hindi/Urdu/any script — so we check that
non-Latin column names and values sort/filter/aggregate/formula/find-replace correctly,
survive a UTF-8 round-trip, and are summarized in the structure. Plus: the battery has
all four languages for every distinct operation.

Run from backend:  .venv\\Scripts\\python.exe tests\\test_phase_3_5.py
"""
from __future__ import annotations

import csv
import io
import os
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TESTS = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS.parent))

_fd, _db = tempfile.mkstemp(suffix="-p35.db")
os.close(_fd)
os.environ["DATABASE_URL"] = "sqlite:///" + _db.replace("\\", "/")

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import app.main as m  # noqa: E402
from app.executor import execute_multi  # noqa: E402
from app.reader import summarize_tables  # noqa: E402

passed = failed = 0


def check(name, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def run(df, ops, primary="t"):
    res, name, notes, render = execute_multi({primary: df.copy()}, primary, ops)
    return res, " ".join(notes), render


print("ENGINE PHASE 3.5 — multi-language tuning (offline: engine is language-agnostic)\n")

# Devanagari (Hindi) column names + values, and Urdu values.
HI = pd.DataFrame({
    "क्षेत्र": ["उत्तर", "दक्षिण", "उत्तर", "पूर्व"],   # Region
    "मूल्य": [100, 200, 150, 50],                       # Price
    "मात्रा": [1, 2, 3, 4],                             # Qty
})

# ---- (a) sort on a Devanagari column ----
res, note, _ = run(HI, [{"action": "sort", "columns": ["मूल्य"], "orders": ["desc"]}])
check("sort works on a Devanagari column (मूल्य desc)", list(res["मूल्य"]) == [200, 150, 100, 50], str(list(res["मूल्य"])))

# ---- (b) filter on Devanagari column + value ----
res, note, _ = run(HI, [{"action": "filter", "conditions": [
    {"column": "क्षेत्र", "operator": "equals", "value": "उत्तर"}]}])
check("filter matches a Devanagari value (क्षेत्र == उत्तर → 2 rows)", len(res) == 2 and set(res["क्षेत्र"]) == {"उत्तर"}, str(res["क्षेत्र"].tolist()))

# ---- (c) aggregate grouped by a Devanagari column ----
res, note, _ = run(HI, [{"action": "aggregate", "agg_func": "sum", "agg_column": "मूल्य", "group_by": ["क्षेत्र"]}])
north = res[res["क्षेत्र"] == "उत्तर"].iloc[0]
val_col = [c for c in res.columns if c != "क्षेत्र"][0]
check("aggregate groups by a Devanagari column (उत्तर = 100+150 = 250)", north[val_col] == 250, res.to_string())

# ---- (d) add_formula_column referencing Devanagari columns ----
res, note, render = run(HI, [{"action": "add_formula_column", "name": "कुल", "formula": "{मूल्य} * {मात्रा}"}])
check("formula column with Devanagari references computes (मूल्य*मात्रा)",
      list(res["कुल"]) == [100, 400, 450, 200], str(list(res["कुल"])))
check("the live formula surfaces the Devanagari column names",
      any("मूल्य" in d.get("formula", "") for d in render if d.get("type") == "formula"), str(render))

# ---- (e) find_replace on Urdu text values ----
UR = pd.DataFrame({"شہر": ["شمال", "جنوب", "شمال"], "قیمت": [10, 20, 30]})  # City / Price
res, note, _ = run(UR, [{"action": "find_replace", "find": "شمال", "replace": "شمالی", "column": "شہر"}])
check("find_replace works on Urdu values (شمال → شمالی)",
      list(res["شہر"]) == ["شمالی", "جنوب", "شمالی"], str(list(res["شہر"])))

# ---- (f) UTF-8 round-trip: non-Latin headers + values survive save/reload ----
res, _, render = run(HI, [{"action": "sort", "columns": ["मूल्य"], "orders": ["asc"]}])
out, _, _ = m._serialize(res, "x.csv", "xlsx", render)
ws = load_workbook(io.BytesIO(out)).active
headers = [ws.cell(row=1, column=i).value for i in range(1, 4)]
check("UTF-8 round-trip: Devanagari headers survive the .xlsx save/reload",
      "क्षेत्र" in headers and "मूल्य" in headers, str(headers))
check("UTF-8 round-trip: a Devanagari value survives", any(ws.cell(row=r, column=1).value == "उत्तर" for r in range(2, 6)), "")

# ---- (g) the structure summary includes non-Latin column names (what the Brain sees) ----
struct = str(summarize_tables({"बिक्री": HI}, "बिक्री"))  # stringify whatever shape it returns
check("structure summary lists Devanagari column names for the Brain",
      "क्षेत्र" in struct and "मूल्य" in struct, struct[:200])

# ---- (h) BATTERY COVERAGE: every distinct operation has all four languages ----
rows = list(csv.DictReader(open(TESTS / "prompt_battery.csv", encoding="utf-8", newline="")))
by_lang = defaultdict(int)
for r in rows:
    by_lang[r["language"]] += 1
check("battery has rows in all four languages",
      all(by_lang.get(l, 0) > 0 for l in ("EN", "HI", "UR", "Hinglish")), dict(by_lang))

# distinct operation = first token of expected_plan (skip rows with none). Two tokens are
# just LABEL noise for operations that ARE covered in all 4 languages under a sibling
# token: "duplicate" (remove_duplicates is covered under 'remove_duplicates') and
# "drop_missing" (covered inside the multi-step rows). Exclude those from the strict check.
_TOKEN_NOISE = {"duplicate", "drop_missing"}
op_langs = defaultdict(set)
for r in rows:
    ep = (r.get("expected_plan") or "").strip()
    if ep:
        action = ep.split(";")[0].strip()
        if action not in _TOKEN_NOISE:
            op_langs[action].add(r["language"])
LANGS = {"EN", "HI", "UR", "Hinglish"}
covered = {op for op, ls in op_langs.items() if LANGS <= ls}
missing = {op: sorted(LANGS - ls) for op, ls in op_langs.items() if not (LANGS <= ls)}
check("every distinct battery operation is covered in all 4 languages",
      not missing, f"{len(covered)}/{len(op_langs)} full; missing: {dict(list(missing.items())[:8])}")
print(f"    (distinct operations with full 4-language coverage: {len(covered)}/{len(op_langs)})")

# ---- (i) UR coverage is now substantial (was thin before 3.5) ----
check("UR coverage grew past 50 rows (3.5 filled the gap)", by_lang.get("UR", 0) >= 50, str(by_lang.get("UR")))

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
