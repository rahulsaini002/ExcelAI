"""Tests for the dashboard operation (Phase 2.3 — one-page dashboard sheet).

Covers: KPIs compute correctly, charts validated, summary + title carried, data
unchanged, layout doesn't overlap, the Dashboard sheet is really written into the
.xlsx, and it regenerates cleanly when data changes. Run: python test_dashboard.py
"""
import io
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook

from app.executor import execute_multi, OperationError
from app.operations.dashboard import dashboard
import app.main as m

passed = failed = 0


def check(name, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


print("Running dashboard checks...\n")

df = pd.DataFrame({
    "Month": ["Jan", "Feb", "Mar"],
    "Revenue": [100, 200, 150],
    "Orders": [3, 5, 4],
})

spec = {
    "action": "dashboard",
    "dashboard_title": "Shop overview",
    "kpis": [
        {"label": "Total Revenue", "agg": "sum", "column": "Revenue", "format": "currency"},
        {"label": "Orders", "agg": "sum", "column": "Orders"},
        {"label": "Months", "agg": "count"},
    ],
    "charts": [
        {"chart_type": "bar", "x_column": "Month", "y_columns": ["Revenue"], "title": "Revenue by Month"},
        {"chart_type": "line", "x_column": "Month", "y_columns": ["Orders"], "title": "Orders by Month"},
    ],
    "summary": "Revenue peaked in February.",
}

# --- happy: KPIs compute correctly ---
note, d = dashboard(df, spec)
kpis = {k["label"]: k["value"] for k in d["kpis"]}
check("KPI sum as currency", kpis["Total Revenue"] == "₹450", str(kpis))
check("KPI sum as number", kpis["Orders"] == "12", str(kpis))
check("KPI plain row count", kpis["Months"] == "3", str(kpis))
check("two charts validated", len(d["charts"]) == 2, str(d["charts"]))
check("summary carried", d["summary"] == "Revenue peaked in February.", str(d))
check("title carried", d["title"] == "Shop overview", str(d))

# --- data unchanged + flows through the executor ---
res, name, notes, render = execute_multi({"t": df}, "t", [spec])
check("dashboard leaves data unchanged",
      len(res) == 3 and list(res.columns) == ["Month", "Revenue", "Orders"], str(list(res.columns)))
check("dashboard directive in render ops",
      any(r.get("type") == "dashboard" for r in render), str(render))

# --- failures ---
try:
    dashboard(df.iloc[0:0], spec)
    check("empty data errors", False, "no error")
except OperationError:
    check("empty data errors", True)

try:
    dashboard(df, {"action": "dashboard", "kpis": [{"label": "x", "agg": "sum", "column": "Nope"}]})
    check("bad KPI column errors", False, "no error")
except OperationError:
    check("bad KPI column errors", True)

try:
    dashboard(df, {"action": "dashboard",
                   "charts": [{"chart_type": "radar", "x_column": "Month", "y_columns": ["Revenue"]}]})
    check("unsupported dashboard chart errors", False, "no error")
except OperationError:
    check("unsupported dashboard chart errors", True)

try:
    dashboard(df, {"action": "dashboard"})  # no kpis, no charts
    check("empty dashboard errors", False, "no error")
except OperationError:
    check("empty dashboard errors", True)

# --- layout written to a Dashboard sheet; charts don't overlap the text block ---
wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(list(df.columns))
for r in df.itertuples(index=False):
    ws.append(list(r))


class _FakeWriter:  # mimics what _apply_dashboard reads from a pandas ExcelWriter
    pass


writer = _FakeWriter()
writer.book = wb
writer.sheets = {"Sheet1": ws}
m._apply_dashboard(writer, "Sheet1", df, d)
dash = wb["Dashboard"]
check("Dashboard sheet created", "Dashboard" in wb.sheetnames)
check("title in A1", dash["A1"].value == "Shop overview", str(dash["A1"].value))
check("first KPI value in B4", dash["B4"].value == "₹450", str(dash["B4"].value))
check("both charts placed", len(dash._charts) == 2, str(len(dash._charts)))
anchors = [str(c.anchor) for c in dash._charts]
check("charts at distinct anchors (no overlap)", len(set(anchors)) == 2, str(anchors))
check("charts sit right of the A–C text block (col E)",
      all(a.startswith("E") for a in anchors), str(anchors))

# --- end-to-end serialize: .xlsx has a Dashboard sheet + chart parts ---
out, fname, media = m._serialize(df, "shop.csv", "xlsx", [d])
wb2 = load_workbook(io.BytesIO(out))
check("serialized .xlsx has a Dashboard sheet", "Dashboard" in wb2.sheetnames, str(wb2.sheetnames))
parts = zipfile.ZipFile(io.BytesIO(out)).namelist()
check("serialized .xlsx contains chart parts", any("chart" in p for p in parts), "")

# --- regenerates cleanly when the data changes ---
df2 = pd.DataFrame({"Month": ["Jan", "Feb"], "Revenue": [1000, 2000], "Orders": [10, 20]})
_, d2 = dashboard(df2, spec)
kpis2 = {k["label"]: k["value"] for k in d2["kpis"]}
check("KPIs reflect changed data", kpis2["Total Revenue"] == "₹3.0K", str(kpis2))
out2, _, _ = m._serialize(df2, "shop.csv", "xlsx", [d2])
wb3 = load_workbook(io.BytesIO(out2))
check("exactly one Dashboard sheet after regen", wb3.sheetnames.count("Dashboard") == 1, str(wb3.sheetnames))

print(f"\n{passed} passed, {failed} failed.")
raise SystemExit(1 if failed else 0)
