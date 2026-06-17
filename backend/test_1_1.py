"""Feature 1.1 — File upload & structure detection. Verifies every acceptance
criterion + test case (1.1-a..k) plus edge cases, at the reader level (no API key).

Run from backend:  .venv\\Scripts\\python.exe test_1_1.py
"""
from __future__ import annotations

import io
import json

import openpyxl
import pandas as pd

from app.reader import load_files, summarize_structure

passed = 0
failed = 0


def check(name, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def load(filename, data):
    return load_files([(filename, data)])


print("FEATURE 1.1 — upload & structure detection\n")

# 1.1-a Standard single-sheet xlsx, 5 columns
b = xlsx_bytes({"Sheet1": [["Date", "Item", "Qty", "Price", "Total"], ["2021-01-01", "Pen", 2, 10, 20]]})
d = load("sales.xlsx", b)
t = next(iter(d.tables.values()))
s = summarize_structure(t)
check("1.1-a columns read", [c["name"] for c in s["columns"]] == ["Date", "Item", "Qty", "Price", "Total"], str(s["columns"]))
check("1.1-a row count", s["row_count"] == 1)
check("1.1-a preview present", len(s["sample_rows"]) == 1)

# 1.1-b CSV, 3 columns -> one table
d = load("data.csv", b"A,B,C\n1,2,3\n4,5,6\n")
check("1.1-b csv single table", len(d.tables) == 1 and list(d.tables)[0] == "data")
check("1.1-b csv columns", [c["name"] for c in summarize_structure(d.tables["data"])["columns"]] == ["A", "B", "C"])

# 1.1-c Multi-sheet workbook -> all sheets listed
b = xlsx_bytes({
    "Jan": [["x", "y"], [1, 2]],
    "Feb": [["p", "q"], [3, 4]],
    "Mar": [["m"], [5]],
})
d = load("book.xlsx", b)
check("1.1-c three sheets listed", set(d.tables) == {"book - Jan", "book - Feb", "book - Mar"}, str(list(d.tables)))
check("1.1-c each sheet own columns", [c["name"] for c in summarize_structure(d.tables["book - Feb"])["columns"]] == ["p", "q"])

# 1.1-d Blank header -> placeholder "Column X"
b = xlsx_bytes({"Sheet1": [["Name", None, "Qty"], ["A", "keep", 1], ["B", "keep2", 2]]})
d = load("blank.xlsx", b)
cols = [c["name"] for c in summarize_structure(next(iter(d.tables.values())))["columns"]]
check("1.1-d blank header gets placeholder", "Name" in cols and "Qty" in cols and any(c.startswith("Column ") for c in cols) and not any("Unnamed" in c for c in cols), str(cols))

# 1.1-e Duplicate headers -> disambiguated
d = load("dup.csv", b"Amount,Amount\n1,2\n3,4\n")
cols = [c["name"] for c in summarize_structure(d.tables["dup"])["columns"]]
check("1.1-e duplicate headers disambiguated", cols == ["Amount", "Amount.1"], str(cols))

# 1.1-f Mixed data column -> text
d = load("mixed.csv", "v\n1\ntwo\n3\n".encode())
typ = summarize_structure(d.tables["mixed"])["columns"][0]["type"]
check("1.1-f mixed column is text", typ == "text", typ)

# basic type detection: number + date
d = load("types.csv", "n,when\n1.5,2021-01-01\n2.5,2022-06-01\n".encode())
types = {c["name"]: c["type"] for c in summarize_structure(d.tables["types"])["columns"]}
check("type: number detected", types["n"] == "number", str(types))
check("type: date detected", types["when"] == "date", str(types))

# 1.1-g Empty file (header only, no rows)
d = load("empty.csv", b"A,B\n")
check("1.1-g empty -> 0 rows, no crash", summarize_structure(d.tables["empty"])["row_count"] == 0)

# 1.1-h Wrong file type
try:
    load("pic.pdf", b"%PDF-1.4 not a sheet")
    check("1.1-h wrong type rejected", False, "no error")
except ValueError as e:
    check("1.1-h wrong type message", "Excel" in str(e) and "CSV" in str(e), str(e))

# 1.1-i Corrupted xlsx
try:
    load("broken.xlsx", b"this is definitely not a real xlsx")
    check("1.1-i corrupted rejected", False, "no error")
except ValueError as e:
    check("1.1-i corrupted friendly", "Couldn't read" in str(e) and "re-saving" in str(e).lower() or "re-saving" in str(e), str(e))

# 1.1-j Very large file -> loads, preview limited to 5
big = "id,val\n" + "\n".join(f"{i},{i*2}" for i in range(100000))
d = load("big.csv", big.encode())
s = summarize_structure(d.tables["big"])
check("1.1-j large file loads", s["row_count"] == 100000)
check("1.1-j preview limited to 5", len(s["sample_rows"]) == 5, str(len(s["sample_rows"])))

# 1.1-k Non-English (Hindi) headers preserved (UTF-8)
d = load("hin.csv", "नाम,मूल्य\nराम,१०\n".encode("utf-8"))
cols = [c["name"] for c in summarize_structure(d.tables["hin"])["columns"]]
check("1.1-k non-English headers preserved", cols == ["नाम", "मूल्य"], str(cols))

# .xlsm accepted
b = xlsx_bytes({"Sheet1": [["a", "b"], [1, 2]]})
try:
    d = load("macro.xlsm", b)  # same zip format; pandas reads it
    check(".xlsm accepted", len(d.tables) == 1)
except ValueError as e:
    check(".xlsm accepted", False, str(e))

# JSON-safety: summarize output (incl. a date column) must be JSON-serializable
d = load("dt.csv", "when\n2021-01-01\n2022-02-02\n".encode())
try:
    json.dumps(summarize_structure(d.tables["dt"]))
    check("structure is JSON-serializable (dates safe)", True)
except TypeError as e:
    check("structure is JSON-serializable (dates safe)", False, str(e))

print(f"\n{passed} passed, {failed} failed.")
raise SystemExit(1 if failed else 0)
